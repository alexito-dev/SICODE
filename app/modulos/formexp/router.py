from fastapi import APIRouter, Request, Depends
from fastapi.responses import HTMLResponse, JSONResponse
from fastapi.templating import Jinja2Templates
from pydantic import BaseModel, field_validator
import openpyxl
import os
from app.modulos.auth.dependencies import login_required
from app.core.config import settings

# xlwings opcional
try:
    import xlwings as xw
    HAS_XLWINGS = True
except ImportError:
    HAS_XLWINGS = False

router = APIRouter(
    prefix="/formexp",
    tags=["formexp"],
    dependencies=[Depends(login_required)]
)

templates_dir = settings.get_template_path("app/modulos/formexp/templates")
templates_base = settings.get_template_path("app/modulos/inicio/templates")
templates = Jinja2Templates(directory=[templates_dir, templates_base])

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
EXCEL_PATH = os.path.join(BASE_DIR, "FORMULARIO194_MESADEPARTES.xlsx")

HEADERS = [
    'N° CAJA', 'N° PAQUETE', 'NUMERO DE REGISTRO', 'TOMO',
    'RANGO INICIAL', 'RANGO FINAL', 'FOLIOS', 'TIPO DOCUMENTAL',
    'N° DE DOCUMENTO', 'RAZON SOCIAL', 'RUC', 'FECHA EXTREMA',
    'OBSERVACIONES', 'X1', 'X2', 'X3'
]

# ─── MODELS ────────────────────────────────────────────
class RegistroFormulario(BaseModel):
    documento: str
    razon_social: str = ""
    ruc: str
    fecha: str = ""
    observaciones: str = ""

    @field_validator("ruc")
    @classmethod
    def ruc_must_be_11(cls, v):
        v = v.strip()
        if len(v) != 11 or not v.isdigit():
            raise ValueError("RUC debe tener exactamente 11 dígitos")
        return v

class RegistroExpediente(BaseModel):
    parte1: str
    parte2: str
    parte3: str
    parte4: str
    razon_social: str = ""
    ruc: str
    fecha: str = ""
    observaciones: str = ""

    @field_validator("ruc")
    @classmethod
    def ruc_must_be_11(cls, v):
        v = v.strip()
        if len(v) != 11 or not v.isdigit():
            raise ValueError("RUC debe tener exactamente 11 dígitos")
        return v

# ─── EXCEL HELPERS ─────────────────────────────────────
def ensure_excel():
    if os.path.exists(EXCEL_PATH):
        return
    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "FORMULARIO 194"
    for i, h in enumerate(HEADERS, 1):
        ws1.cell(1, i, h)
    ws2 = wb.create_sheet("EXPEDIENTES")
    for i, h in enumerate(HEADERS, 1):
        ws2.cell(1, i, h)
    wb.save(EXCEL_PATH)

def get_xlwings_book():
    if not HAS_XLWINGS:
        return None
    try:
        fname = os.path.basename(EXCEL_PATH)
        for a in xw.apps:
            for b in a.books:
                if b.name.lower() == fname.lower():
                    return b
    except Exception:
        pass
    return None

def write_record(sheet_name: str, data: dict) -> dict:
    """Escribe un registro. Intenta xlwings primero, luego openpyxl."""
    live = False
    book = get_xlwings_book()

    if book:
        try:
            ws = book.sheets[sheet_name]
            if not ws.range("H2").value:
                row = 2
            else:
                row = ws.range("H1").end("down").row + 1
            for col, val in data.items():
                ws.range((row, col)).value = val
            live = True
        except Exception:
            live = False

    if not live:
        wb = openpyxl.load_workbook(EXCEL_PATH)
        ws = wb[sheet_name]
        row = ws.max_row + 1
        for col, val in data.items():
            ws.cell(row, col, val)
        wb.save(EXCEL_PATH)
        wb.close()

    return {"success": True, "live": live, "row": row}

def count_records(sheet_name: str) -> int:
    book = get_xlwings_book()
    if book:
        try:
            ws = book.sheets[sheet_name]
            if not ws.range("H2").value:
                return 0
            return max(0, ws.range("H1").end("down").row - 1)
        except Exception:
            pass
    try:
        wb = openpyxl.load_workbook(EXCEL_PATH, read_only=True)
        ws = wb[sheet_name]
        n = max(0, ws.max_row - 1) if ws.max_row else 0
        wb.close()
        return n
    except Exception:
        return 0

# ─── INIT EXCEL ────────────────────────────────────────
ensure_excel()

# ─── ROUTES ────────────────────────────────────────────
@router.get("/")
async def formexp_home(request: Request):
    return templates.TemplateResponse("formexp/index.html", {
        "request": request,
        "modulo": "FormExp"
    })

@router.get("/api/status")
async def api_status():
    book = get_xlwings_book()
    return {
        "excel_connected": book is not None,
        "has_xlwings": HAS_XLWINGS,
        "formulario_count": count_records("FORMULARIO 194"),
        "expediente_count": count_records("EXPEDIENTES"),
    }

@router.post("/api/formulario")
async def save_formulario(reg: RegistroFormulario):
    fecha_val = reg.fecha.strip()
    if fecha_val:
        fecha_val = f"'{fecha_val}"

    data = {
        8: "FORMULARIO NRO. 194",
        9: reg.documento.strip(),
        10: reg.razon_social.strip(),
        11: reg.ruc.strip(),
        12: fecha_val,
        13: reg.observaciones.strip(),
    }
    result = write_record("FORMULARIO 194", data)
    result["count"] = count_records("FORMULARIO 194")
    return result

@router.post("/api/expediente")
async def save_expediente(reg: RegistroExpediente):
    fecha_val = reg.fecha.strip()
    if fecha_val:
        fecha_val = f"'{fecha_val}"

    doc = f"000-URD{reg.parte1.strip()}-{reg.parte2.strip()}-{reg.parte3.strip()}-{reg.parte4.strip()}"
    data = {
        8: "EXPEDIENTE DE MEZA DE PARTES",
        9: doc,
        10: reg.razon_social.strip(),
        11: reg.ruc.strip(),
        12: fecha_val,
        13: reg.observaciones.strip(),
    }
    result = write_record("EXPEDIENTES", data)
    result["count"] = count_records("EXPEDIENTES")
    return result
