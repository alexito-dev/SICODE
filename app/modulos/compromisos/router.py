"""
Módulo Compromisos - Registro de llamadas a contribuyentes y compromisos de pago,
con listado filtrable, resumen y exportación a Excel.
"""

import io
from datetime import date, datetime
from typing import Optional

from fastapi import APIRouter, Request, Depends, HTTPException
from fastapi.responses import StreamingResponse
from fastapi.templating import Jinja2Templates
from pydantic import BaseModel
from sqlalchemy import or_, desc

from app.core.config import settings
from app.core.database import SessionLocal
from app.modulos.auth.dependencies import login_required
from app.modulos.compromisos.models import Llamada, RESULTADOS, ESTADOS

# Textos repetidos
_PENDIENTE = "Pendiente"
_STATUS = "status"
_SUCCESS = "success"

router = APIRouter(prefix="/compromisos", tags=["compromisos"], dependencies=[Depends(login_required)])

templates_dir = settings.get_template_path("app/modulos/compromisos/templates")
templates_base = settings.get_template_path("app/modulos/inicio/templates")
templates = Jinja2Templates(directory=[templates_dir, templates_base])


class LlamadaIn(BaseModel):
    ruc: str
    razon_social: str = ""
    telefono: str = ""
    contacto: str = ""
    fecha_llamada: Optional[str] = None
    resultado: str = "Contestó"
    fecha_compromiso: Optional[str] = None
    monto_compromiso: Optional[float] = None
    estado: str = _PENDIENTE
    proxima_llamada: Optional[str] = None
    observaciones: str = ""


def _fecha(valor: Optional[str]) -> Optional[date]:
    if not valor:
        return None
    try:
        return datetime.strptime(valor[:10], "%Y-%m-%d").date()
    except ValueError:
        raise HTTPException(status_code=400, detail=f"Fecha inválida: {valor}") from None


def _aplicar(obj: Llamada, d: LlamadaIn):
    ruc = d.ruc.strip()
    if not ruc.isdigit() or len(ruc) != 11:
        raise HTTPException(status_code=400, detail="El RUC debe tener 11 dígitos.")
    if d.resultado not in RESULTADOS or d.estado not in ESTADOS:
        raise HTTPException(status_code=400, detail="Resultado o estado no válido.")
    obj.ruc = ruc
    obj.razon_social = d.razon_social.strip()
    obj.telefono = d.telefono.strip()
    obj.contacto = d.contacto.strip()
    if d.fecha_llamada:
        try:
            obj.fecha_llamada = datetime.strptime(d.fecha_llamada[:16], "%Y-%m-%dT%H:%M")
        except ValueError:
            raise HTTPException(status_code=400, detail="Fecha de llamada inválida.") from None
    obj.resultado = d.resultado
    obj.fecha_compromiso = _fecha(d.fecha_compromiso)
    obj.monto_compromiso = d.monto_compromiso
    obj.estado = d.estado
    obj.proxima_llamada = _fecha(d.proxima_llamada)
    obj.observaciones = d.observaciones.strip()


def _consulta(db, q, resultado, estado, desde, hasta):
    query = db.query(Llamada).filter(Llamada.is_active)
    if q and q.strip():
        t = f"%{q.strip()}%"
        query = query.filter(
            or_(
                Llamada.ruc.ilike(t),
                Llamada.razon_social.ilike(t),
                Llamada.telefono.ilike(t),
                Llamada.observaciones.ilike(t),
            )
        )
    if resultado:
        query = query.filter(Llamada.resultado == resultado)
    d, h = _fecha(desde), _fecha(hasta)
    if d:
        query = query.filter(Llamada.fecha_llamada >= datetime.combine(d, datetime.min.time()))
    if h:
        query = query.filter(Llamada.fecha_llamada <= datetime.combine(h, datetime.max.time()))
    filas = query.order_by(desc(Llamada.fecha_llamada)).all()
    if estado:
        filas = [f for f in filas if f.estado_efectivo == estado]
    return filas


@router.get("/")
async def home(request: Request):
    return templates.TemplateResponse(
        "compromisos/index.html",
        {"request": request, "title": "Compromisos de Pago", "resultados": RESULTADOS, "estados": ESTADOS},
    )


def _monto(filas) -> float:
    return sum(f.monto_compromiso or 0 for f in filas)


def _contar(filas, condicion) -> int:
    return sum(1 for f in filas if condicion(f))


def _resumen(filas, hoy) -> dict:
    """Totales del tablero: llamadas, compromisos y montos por estado."""
    con_comp = [f for f in filas if f.fecha_compromiso]
    return {
        "llamadas": len(filas),
        "contestadas": _contar(filas, lambda f: f.resultado == "Contestó"),
        "compromisos": len(con_comp),
        "monto_total": _monto(con_comp),
        "monto_pendiente": _monto([f for f in con_comp if f.estado_efectivo == _PENDIENTE]),
        "monto_vencido": _monto([f for f in con_comp if f.estado_efectivo == "Vencido"]),
        "monto_cumplido": _monto([f for f in con_comp if f.estado == "Cumplido"]),
        "vencen_hoy": _contar(con_comp, lambda f: f.estado == _PENDIENTE and f.fecha_compromiso == hoy),
        "por_llamar": _contar(filas, lambda f: bool(f.proxima_llamada) and f.proxima_llamada <= hoy),
    }


@router.get("/api/llamadas")
async def listar(
    q: Optional[str] = None,
    resultado: Optional[str] = None,
    estado: Optional[str] = None,
    desde: Optional[str] = None,
    hasta: Optional[str] = None,
):
    db = SessionLocal()
    try:
        filas = _consulta(db, q, resultado, estado, desde, hasta)
        resumen = _resumen(filas, date.today())
        return {_STATUS: _SUCCESS, "resumen": resumen, "data": [f.to_dict() for f in filas]}
    finally:
        db.close()


@router.post("/api/llamadas")
async def crear(request: Request, datos: LlamadaIn):
    db = SessionLocal()
    try:
        obj = Llamada()
        _aplicar(obj, datos)
        user = getattr(request.state, "user", None)
        obj.usuario_id = user.id if user else None
        obj.usuario_nombre = (getattr(user, "nombre", "") if user else "") or ""
        if not obj.fecha_llamada:
            obj.fecha_llamada = datetime.now()
        db.add(obj)
        db.commit()
        return {_STATUS: _SUCCESS, "data": obj.to_dict()}
    finally:
        db.close()


@router.put("/api/llamadas/{llamada_id}")
async def actualizar(llamada_id: int, datos: LlamadaIn):
    db = SessionLocal()
    try:
        obj = db.query(Llamada).filter(Llamada.id == llamada_id, Llamada.is_active).first()
        if not obj:
            raise HTTPException(status_code=404, detail="Registro no encontrado.")
        _aplicar(obj, datos)
        db.commit()
        return {_STATUS: _SUCCESS, "data": obj.to_dict()}
    finally:
        db.close()


@router.delete("/api/llamadas/{llamada_id}")
async def eliminar(llamada_id: int):
    db = SessionLocal()
    try:
        obj = db.query(Llamada).filter(Llamada.id == llamada_id, Llamada.is_active).first()
        if not obj:
            raise HTTPException(status_code=404, detail="Registro no encontrado.")
        obj.is_active = False
        db.commit()
        return {_STATUS: _SUCCESS}
    finally:
        db.close()


@router.get("/api/exportar")
async def exportar(
    q: Optional[str] = None,
    resultado: Optional[str] = None,
    estado: Optional[str] = None,
    desde: Optional[str] = None,
    hasta: Optional[str] = None,
):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    db = SessionLocal()
    try:
        filas = _consulta(db, q, resultado, estado, desde, hasta)
        wb = Workbook()
        ws = wb.active
        ws.title = "Compromisos"
        ws.append(
            [
                "Fecha llamada",
                "RUC",
                "Razón social",
                "Teléfono",
                "Contacto",
                "Resultado",
                "Fecha compromiso",
                "Monto",
                "Estado",
                "Próxima llamada",
                "Observaciones",
                "Registrado por",
            ]
        )
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="1E3A8A")
            c.alignment = Alignment(horizontal="center")
        dmy = lambda d, f="%d/%m/%Y": d.strftime(f) if d else ""
        for f in filas:
            ws.append(
                [
                    dmy(f.fecha_llamada, "%d/%m/%Y %H:%M"),
                    f.ruc,
                    f.razon_social,
                    f.telefono,
                    f.contacto,
                    f.resultado,
                    dmy(f.fecha_compromiso),
                    f.monto_compromiso,
                    f.estado_efectivo,
                    dmy(f.proxima_llamada),
                    f.observaciones,
                    f.usuario_nombre,
                ]
            )
        for fila in ws.iter_rows(min_row=2, min_col=8, max_col=8):
            fila[0].number_format = "#,##0.00"
        for i, w in enumerate([17, 13, 38, 14, 22, 15, 16, 14, 14, 16, 50, 20], start=1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        return StreamingResponse(
            buf,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f'attachment; filename="Compromisos_{date.today():%Y%m%d}.xlsx"'},
        )
    finally:
        db.close()
