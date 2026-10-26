import os
import math
from fastapi import APIRouter, Request, UploadFile, File, Form
from fastapi.templating import Jinja2Templates
from fastapi.responses import JSONResponse
import fitz  # PyMuPDF
from PIL import Image, ImageOps
import io
import base64
import json

router = APIRouter(prefix="/datamask", tags=["datamask"])

from app.core.config import settings

# Textos repetidos
_ALETS = "alets"
_DESKTOP = "Desktop"
_ERROR = "error"
_ESCRITORIO = "Escritorio"
_ONEDRIVE = "OneDrive"
_STATUS = "status"
_SUCCESS = "success"
_USERS = "Users"

templates_dir = settings.get_template_path("app/modulos/datamask/templates")
templates_base = settings.get_template_path("app/modulos/inicio/templates")
templates = Jinja2Templates(directory=[templates_dir, templates_base])


@router.get("/")
async def datamask_home(request: Request):
    return templates.TemplateResponse("datamask/index.html", {"request": request})


# ─── Process (annotate/mask/images/text) ───────────────────────────────────────


def safe_float(val, default=0.0):
    if val is None:
        return default
    try:
        f = float(val)
        if math.isnan(f) or math.isinf(f):
            return default
        return f
    except (ValueError, TypeError):
        return default


def _parse_json(raw, default):
    try:
        return json.loads(raw or default)
    except Exception:
        return {}


def _rotate_pages(doc, user_page_rotations):
    if not (user_page_rotations and isinstance(user_page_rotations, dict)):
        return
    for p_idx_s, add_rot in user_page_rotations.items():
        try:
            p_idx = int(p_idx_s)
            if 0 <= p_idx < len(doc):
                doc[p_idx].set_rotation((doc[p_idx].rotation + int(add_rot)) % 360)
        except Exception:
            pass


def _ann_color(ann):
    hex_c = ann.get("color", "#000000")
    if isinstance(hex_c, str) and hex_c.startswith("#") and len(hex_c) == 7:
        try:
            return (int(hex_c[1:3], 16) / 255.0, int(hex_c[3:5], 16) / 255.0, int(hex_c[5:7], 16) / 255.0)
        except ValueError:
            pass
    return (0.0, 0.0, 0.0)


class _PageCtx:
    """Página a editar con sus medidas y el estado de lo que se dibujó."""

    def __init__(self, page):
        self.page = page
        self.pw, self.ph = page.rect.width, page.rect.height
        self.page_rot = page.rotation % 360
        self.derot = page.derotation_matrix
        self.shape = page.new_shape()
        self.has_shape_drawings = False
        self.has_redactions = False

    def abs_x(self, value):
        return self.page.rect.x0 + value * self.pw

    def abs_y(self, value):
        return self.page.rect.y0 + value * self.ph

    def point(self, x, y):
        return fitz.Point(self.abs_x(x), self.abs_y(y)) * self.derot

    def rect_from_list(self, ann):
        r_list = ann.get("rect") or [0.1, 0.1, 0.2, 0.1]
        rx = safe_float(r_list[0] if len(r_list) > 0 else 0.1)
        ry = safe_float(r_list[1] if len(r_list) > 1 else 0.1)
        rw = safe_float(r_list[2] if len(r_list) > 2 else 0.2)
        rh = safe_float(r_list[3] if len(r_list) > 3 else 0.1)
        r_vis = fitz.Rect(self.abs_x(rx), self.abs_y(ry), self.abs_x(rx + rw), self.abs_y(ry + rh))
        return (r_vis * self.derot).normalize()


def _apply_highlight(ctx, ann, col):
    ctx.shape.draw_rect(ctx.rect_from_list(ann))
    ctx.shape.finish(fill=col, fill_opacity=0.35)
    ctx.has_shape_drawings = True


def _apply_redact(ctx, ann, col):
    ctx.page.add_redact_annot(ctx.rect_from_list(ann), fill=col)
    ctx.has_redactions = True


def _apply_strike(ctx, ann, col):
    p1_dict = ann.get("p1") if isinstance(ann.get("p1"), dict) else {"x": 0.1, "y": 0.1}
    p2_dict = ann.get("p2") if isinstance(ann.get("p2"), dict) else {"x": 0.2, "y": 0.2}
    pt1 = ctx.point(safe_float(p1_dict.get("x"), 0.1), safe_float(p1_dict.get("y"), 0.1))
    pt2 = ctx.point(safe_float(p2_dict.get("x"), 0.2), safe_float(p2_dict.get("y"), 0.2))
    ctx.shape.draw_line(pt1, pt2)
    ctx.shape.finish(color=col, width=safe_float(ann.get("size"), 2))
    ctx.has_shape_drawings = True


def _apply_ink(ctx, ann, col):
    pts = ann.get("points") if isinstance(ann.get("points"), list) else []
    if len(pts) <= 1:
        return
    valid_pts = [
        ctx.point(safe_float(p.get("x"), 0.1), safe_float(p.get("y"), 0.1)) for p in pts if isinstance(p, dict)
    ]
    if len(valid_pts) > 1:
        ctx.shape.draw_polyline(valid_pts)
        ctx.shape.finish(color=col, width=safe_float(ann.get("size"), 3))
        ctx.has_shape_drawings = True


def _apply_text(ctx, ann, col):
    tx = ctx.abs_x(safe_float(ann.get("x"), 0.1))
    ty = ctx.abs_y(safe_float(ann.get("y"), 0.1))
    tw = safe_float(ann.get("w"), 0.35) * ctx.pw
    th = safe_float(ann.get("h"), 0.08) * ctx.ph
    f_size = safe_float(ann.get("fontSize"), 13)
    total_rot = int((ctx.page_rot + safe_float(ann.get("rotation"), 0)) % 360)

    text_str = str(ann.get("text") or "").strip()
    if not text_str:
        return
    box_w = max(tw, 40.0)
    box_h = max(th, 15.0)
    r_vis = fitz.Rect(tx, ty, tx + box_w, ty + box_h)
    t_rect = (r_vis * ctx.derot).normalize()
    rc = ctx.page.insert_textbox(t_rect, text_str, fontsize=f_size, color=col, rotate=total_rot)
    if rc < 0:
        r_vis = fitz.Rect(tx, ty, tx + box_w + 200.0, ty + box_h + 120.0)
        t_rect = (r_vis * ctx.derot).normalize()
        ctx.page.insert_textbox(t_rect, text_str, fontsize=max(8.0, f_size - 2), color=col, rotate=total_rot)


def _decode_png(img_data_str):
    if "," in img_data_str:
        img_data_str = img_data_str.split(",", 1)[1]
    img_obj = Image.open(io.BytesIO(base64.b64decode(img_data_str)))
    img_obj = ImageOps.exif_transpose(img_obj)
    if img_obj.mode != "RGBA":
        img_obj = img_obj.convert("RGBA")
    img_byte_arr = io.BytesIO()
    img_obj.save(img_byte_arr, format="PNG")
    return img_byte_arr.getvalue()


def _image_rect(ctx, ann, total_rot):
    ix = ctx.abs_x(safe_float(ann.get("x"), 0.1))
    iy = ctx.abs_y(safe_float(ann.get("y"), 0.1))
    iw = max(10.0, safe_float(ann.get("w"), 0.3) * ctx.pw)
    ih = max(10.0, safe_float(ann.get("h"), 0.15) * ctx.ph)
    if total_rot in [90, 270]:
        cx, cy = ix + iw / 2.0, iy + ih / 2.0
        r_vis = fitz.Rect(cx - ih / 2.0, cy - iw / 2.0, cx + ih / 2.0, cy + iw / 2.0)
    else:
        r_vis = fitz.Rect(ix, iy, ix + iw, iy + ih)
    return (r_vis * ctx.derot).normalize()


def _apply_image(ctx, ann, page_idx):
    img_data_str = str(ann.get("image") or "")
    total_rot = int((ctx.page_rot + safe_float(ann.get("rotation"), 0)) % 360)
    if not img_data_str:
        return
    try:
        png_bytes = _decode_png(img_data_str)
        ctx.page.insert_image(
            _image_rect(ctx, ann, total_rot),
            stream=png_bytes,
            rotate=total_rot,
            keep_proportion=False,
            overlay=True,
        )
    except Exception as img_err:
        print(f"Error inserting image on page {page_idx}: {img_err}")


_ANN_HANDLERS = {
    "highlight": _apply_highlight,
    "redact": _apply_redact,
    "strike": _apply_strike,
    "ink": _apply_ink,
    "text": _apply_text,
}


def _apply_annotation(ctx, ann, page_idx):
    col = _ann_color(ann)
    kind = ann.get("type")
    if kind == "image":
        _apply_image(ctx, ann, page_idx)
        return
    handler = _ANN_HANDLERS.get(kind)
    if handler:
        handler(ctx, ann, col)


def _edit_page(doc, page_idx, anns):
    ctx = _PageCtx(doc[page_idx])
    for ann in anns:
        if not isinstance(ann, dict):
            continue
        try:
            _apply_annotation(ctx, ann, page_idx)
        except Exception as ann_item_err:
            print(f"Error processing ann item on page {page_idx}: {ann_item_err}")

    if ctx.has_shape_drawings:
        try:
            ctx.shape.commit()
        except Exception as shape_err:
            print(f"Shape commit error on page {page_idx}: {shape_err}")

    if ctx.has_redactions:
        try:
            ctx.page.apply_redactions()
        except Exception as red_err:
            print(f"Redaction error on page {page_idx}: {red_err}")


def _apply_annotations(doc, data):
    if not isinstance(data, dict):
        return
    for page_idx_str, anns in data.items():
        try:
            page_idx = int(page_idx_str)
        except ValueError:
            continue
        if page_idx >= len(doc) or not anns or not isinstance(anns, list):
            continue
        _edit_page(doc, page_idx, anns)


def _save_desktop_copy(filename, pdf_bytes):
    """Copia de respaldo en el Escritorio (no interrumpe si falla)."""
    try:
        base_name = os.path.splitext(filename or "documento.pdf")[0] + "_editado.pdf"
        if not base_name.endswith(".pdf"):
            base_name += ".pdf"

        escritorio_onedrive = os.path.join("C:", os.sep, _USERS, _ALETS, _ONEDRIVE, _ESCRITORIO, base_name)
        escritorio_local = os.path.join("C:", os.sep, _USERS, _ALETS, _DESKTOP, base_name)

        for p in [escritorio_onedrive, escritorio_local]:
            try:
                os.makedirs(os.path.dirname(p), exist_ok=True)
                with open(p, "wb") as f_out:
                    f_out.write(pdf_bytes)
                print(f"=== SUCESSFULLY SAVED TO DESKTOP: {p} ===")
            except Exception as ex:
                print(f"Desktop path save warning: {ex}")
    except Exception as dt_err:
        print(f"Desktop save warning: {dt_err}")


@router.post("/process")
async def process_pdf(file: UploadFile = File(...), annotations: str = Form(...), pageRotations: str = Form("{}")):
    try:
        data = _parse_json(annotations, "{}")
        user_page_rotations = _parse_json(pageRotations, "{}")

        pdf_bytes = await file.read()
        doc = fitz.open(stream=pdf_bytes, filetype="pdf")

        _rotate_pages(doc, user_page_rotations)
        _apply_annotations(doc, data)

        buf = io.BytesIO()
        doc.save(buf)
        doc.close()
        buf.seek(0)
        pdf_bytes = buf.getvalue()
        pdf_b64 = base64.b64encode(pdf_bytes).decode()

        _save_desktop_copy(file.filename, pdf_bytes)

        return JSONResponse(content={_STATUS: _SUCCESS, "file": pdf_b64})

    except Exception as e:
        import traceback

        traceback.print_exc()
        return JSONResponse(status_code=200, content={_STATUS: _ERROR, _ERROR: str(e)})


@router.post("/save_to_desktop")
async def save_to_desktop(file: UploadFile = File(...), annotations: str = Form(...), pageRotations: str = Form("{}")):
    try:
        res = await process_pdf(file=file, annotations=annotations, pageRotations=pageRotations)
        body = json.loads(res.body.decode())
        if body.get(_STATUS) == _SUCCESS and body.get("file"):
            pdf_b64 = body["file"]
            pdf_bytes = base64.b64decode(pdf_b64)

            filename = file.filename or "documento.pdf"
            base_name = os.path.splitext(filename)[0] + "_editado.pdf"
            if not base_name.endswith(".pdf"):
                base_name += ".pdf"

            escritorio_onedrive = os.path.join("C:", os.sep, _USERS, _ALETS, _ONEDRIVE, _ESCRITORIO, base_name)
            escritorio_local = os.path.join("C:", os.sep, _USERS, _ALETS, _DESKTOP, base_name)

            paths_to_save = [
                escritorio_onedrive,
                escritorio_local,
                os.path.join(os.path.expanduser("~"), _ONEDRIVE, _ESCRITORIO, base_name),
                os.path.join(os.path.expanduser("~"), _DESKTOP, base_name),
            ]

            saved_paths = []
            for p in paths_to_save:
                try:
                    os.makedirs(os.path.dirname(p), exist_ok=True)
                    with open(p, "wb") as f_out:
                        f_out.write(pdf_bytes)
                    saved_paths.append(p)
                except Exception:
                    pass

            return JSONResponse(
                content={
                    _STATUS: _SUCCESS,
                    "message": f"Guardado en Escritorio como {base_name}",
                    "paths": saved_paths,
                    "file": pdf_b64,
                }
            )
        return res
    except Exception as e:
        return JSONResponse(status_code=200, content={_STATUS: _ERROR, _ERROR: str(e)})


@router.post("/convert")
async def convert_pdf_to_word(
    file: UploadFile = File(...),
    ocr: str = Form("1"),
):
    """
    Convierte PDF a DOCX usando pdf2docx (texto nativo) y
    easyocr con fallback a pdf2docx para páginas escaneadas.
    """
    try:
        pdf_bytes = await file.read()
        use_ocr = ocr == "1"

        # ── Detectar si el PDF tiene texto nativo ──
        src_doc = fitz.open(stream=pdf_bytes, filetype="pdf")
        total_chars = sum(len(src_doc[i].get_text()) for i in range(min(3, len(src_doc))))
        is_scanned = total_chars < 100  # muy poco texto → escaneado

        docx_bytes = None
        method = "native"

        if is_scanned and use_ocr:
            try:
                docx_bytes = _ocr_to_docx(src_doc)
                method = "ocr"
            except Exception as ocr_err:
                print(f"OCR no disponible o falló ({ocr_err}), usando pdf2docx como fallback...")
                docx_bytes = _pdf2docx_convert(pdf_bytes)
                method = "native_fallback"
        else:
            docx_bytes = _pdf2docx_convert(pdf_bytes)

        src_doc.close()
        return JSONResponse(
            content={_STATUS: _SUCCESS, "file": base64.b64encode(docx_bytes).decode(), "method": method}
        )

    except ImportError as e:
        return JSONResponse(
            status_code=500, content={_ERROR: f"Librería requerida: {e}. Ejecuta: pip install pdf2docx"}
        )
    except Exception as e:
        import traceback

        traceback.print_exc()
        return JSONResponse(status_code=200, content={_STATUS: _ERROR, _ERROR: str(e)})


def _pdf2docx_convert(pdf_bytes: bytes) -> bytes:
    """Usa pdf2docx para convertir PDF a DOCX de manera segura en Windows."""
    from pdf2docx import Converter
    import tempfile
    import os

    pdf_fd, pdf_path = tempfile.mkstemp(suffix=".pdf")
    os.close(pdf_fd)
    with open(pdf_path, "wb") as f:
        f.write(pdf_bytes)

    out_path = pdf_path.replace(".pdf", ".docx")
    try:
        cv = Converter(pdf_path)
        cv.convert(out_path, start=0, end=None)
        cv.close()
        with open(out_path, "rb") as f:
            return f.read()
    finally:
        for p in (pdf_path, out_path):
            try:
                if os.path.exists(p):
                    os.unlink(p)
            except Exception:
                pass


def _ocr_lines(results) -> list:
    """Agrupa los textos del OCR en líneas según su posición vertical."""
    results_sorted = sorted(results, key=lambda x: x[0][0][1])  # sort by y
    current_line_y = results_sorted[0][0][0][1]
    lines = []
    line_texts = []
    for bbox, text, _conf in results_sorted:
        y = bbox[0][1]
        if abs(y - current_line_y) > 20:
            lines.append(" ".join(line_texts))
            line_texts = [text]
            current_line_y = y
        else:
            line_texts.append(text)
    lines.append(" ".join(line_texts))
    return lines


def _ocr_page_results(reader, page):
    """Renderiza la página a alta resolución (mejor OCR) y devuelve lo que lee EasyOCR."""
    import numpy as np
    import cv2

    pix = page.get_pixmap(matrix=fitz.Matrix(2.0, 2.0), alpha=False)
    img_arr = np.frombuffer(pix.tobytes("png"), dtype=np.uint8)
    img = cv2.imdecode(img_arr, cv2.IMREAD_COLOR)
    return reader.readtext(img, detail=1, paragraph=False)


def _ocr_to_docx(src_doc) -> bytes:
    """
    Para PDFs escaneados: usa EasyOCR si está disponible para extraer texto
    y genera un documento DOCX.
    """
    import easyocr
    from docx import Document
    from docx.shared import Pt

    reader = easyocr.Reader(["es", "en"], gpu=False)
    out_doc = Document()
    out_doc.add_heading("Documento Convertido (OCR)", 0)

    for page_num in range(len(src_doc)):
        results = _ocr_page_results(reader, src_doc[page_num])

        if page_num > 0:
            out_doc.add_page_break()
        out_doc.add_heading(f"Página {page_num + 1}", level=2)

        lines = _ocr_lines(results) if results else []
        for k, line in enumerate(lines):
            p = out_doc.add_paragraph(line)
            # La última línea de la página conserva el tamaño por defecto
            if k < len(lines) - 1 and p.runs:
                p.runs[0].font.size = Pt(11)

    buf = io.BytesIO()
    out_doc.save(buf)
    buf.seek(0)
    return buf.getvalue()
