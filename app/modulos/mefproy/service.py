"""
MEFProy Service - Conector y automatizador para la plataforma de Proyectos PTE del MEF.
"""

import urllib.request
import ssl
import json
import io
import os
import base64
import subprocess
import logging
from datetime import datetime
from typing import List, Dict, Any

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# Textos repetidos
_CALIBRI = "Calibri"
_CENTER = "center"
_COD_DGPP = "cod_dgpp"
_COD_SNIP = "cod_snip"
_COLOR = "color"
_COSTO_PROYECTO = "costo_proyecto"
_DEPARTAMENTO = "departamento"
_E2E8F0 = "E2E8F0"
_EJECUCION = "ejecucion"
_ICONO = "icono"
_IGNORE = "ignore"
_K_0163AC = "0163AC"
_NACIONAL = "Nacional"
_NOMBRE = "nombre"
_NOM_PROYECTO = "nom_proyecto"
_PCODEJECUTORA = "pCodEjecutora"
_PERIODO = "periodo"
_PPAGENUMBER = "pPageNumber"
_PPAGESIZE = "pPageSize"
_PPERIODO = "pPeriodo"
_PSORTCOLUMN = "pSortColumn"
_PSORTORDER = "pSortOrder"
_RIGHT = "right"
_SECTOR = "sector"
_UTF_8 = "utf-8"

logger = logging.getLogger("MEFProy-Service")

UNIDADES_EJECUTORAS_PRESET = [
    {"id": "818", _NOMBRE: "GOBIERNO REGIONAL JUNIN - SEDE CENTRAL", _DEPARTAMENTO: "Junín"},
    {"id": "728", _NOMBRE: "PROVIAS NACIONAL - MTC", _DEPARTAMENTO: "Lima / Nacional"},
    {"id": "796", _NOMBRE: "MUNICIPALIDAD METROPOLITANA DE LIMA", _DEPARTAMENTO: "Lima"},
    {"id": "803", _NOMBRE: "GOBIERNO REGIONAL AREQUIPA - SEDE CENTRAL", _DEPARTAMENTO: "Arequipa"},
    {"id": "805", _NOMBRE: "GOBIERNO REGIONAL CUSCO - SEDE CENTRAL", _DEPARTAMENTO: "Cusco"},
    {"id": "814", _NOMBRE: "GOBIERNO REGIONAL PIURA - SEDE CENTRAL", _DEPARTAMENTO: "Piura"},
    {"id": "815", _NOMBRE: "GOBIERNO REGIONAL LA LIBERTAD - SEDE CENTRAL", _DEPARTAMENTO: "La Libertad"},
    {"id": "1209", _NOMBRE: "AUTORIDAD NACIONAL DE INFRAESTRUCTURA (ANIN)", _DEPARTAMENTO: _NACIONAL},
    {"id": "570", _NOMBRE: "MINISTERIO DE EDUCACION - SEDE CENTRAL", _DEPARTAMENTO: _NACIONAL},
    {"id": "573", _NOMBRE: "MINISTERIO DE SALUD - SEDE CENTRAL", _DEPARTAMENTO: _NACIONAL},
]

MEF_BASE_URL = "https://ofi5.mef.gob.pe/proyectos_pte/forms/UnidadEjecutora.aspx"


def detectar_sector(nombre: str) -> Dict[str, str]:
    """Clasifica temáticamente el proyecto para enriquecer la visualización."""
    n = nombre.upper()
    if any(k in n for k in ["EDUCACI", "COLEGIO", " I.E", "INSTITUT", "ESCUELA", "AULA"]):
        return {_NOMBRE: "Educación", _ICONO: "fa-graduation-cap", _COLOR: "blue"}
    if any(k in n for k in ["SALUD", "HOSPITAL", "POSTA", "MEDIC", "CENTRO DE SALUD"]):
        return {_NOMBRE: "Salud", _ICONO: "fa-heart-pulse", _COLOR: "red"}
    if any(k in n for k in ["CARRETERA", "PUENTE", "VIAL", "CAMINO", "PISTA", "TRANSPORTE", "ASFALT", "PAVIMENT"]):
        return {_NOMBRE: "Transporte / Vías", _ICONO: "fa-road", _COLOR: "amber"}
    if any(k in n for k in ["AGUA", "SANEAMIENTO", "ALCANTARILLADO", "DESAGUE", "RESERVORIO", "POTABLE"]):
        return {_NOMBRE: "Saneamiento / Agua", _ICONO: "fa-faucet-drip", _COLOR: "cyan"}
    if any(k in n for k in ["RIEGO", "CANAL", "AGRO", "AGRICOLA", "REPRESA", "PRODUCTIV"]):
        return {_NOMBRE: "Riego / Agro", _ICONO: "fa-seedling", _COLOR: "green"}
    if any(k in n for k in ["ELECTRI", "ENERGIA", "REDES ELECTRICAS"]):
        return {_NOMBRE: "Electrificación", _ICONO: "fa-bolt", _COLOR: "yellow"}
    if any(k in n for k in ["SEGURIDAD", "COMISARIA", "MUNICIPAL", "PALACIO", "CIVICO", "INSTITUC"]):
        return {_NOMBRE: "Institucional / Seguridad", _ICONO: "fa-building-columns", _COLOR: "purple"}
    return {_NOMBRE: "Infraestructura General", _ICONO: "fa-helmet-safety", _COLOR: "slate"}


_FICHA_URL = "https://ofi5.mef.gob.pe/invierteweb/Consulta/ConsultaFicha/ConsultaFicha?CodigoUnico={}"


def _num(valor) -> float:
    """Número de una celda de la grilla; vacío = 0.0."""
    return float(valor) if valor else 0.0


def _ficha_url(cod_dgpp: str, cod_snip: str):
    cod_ref = cod_dgpp or cod_snip
    return _FICHA_URL.format(cod_ref) if cod_ref else None


def _fila_ejecucion(row: Dict[str, Any]) -> Dict[str, Any]:
    cells = row.get("cell", [])
    c_dgpp = str(cells[0]).strip()
    c_snip = str(cells[1]).strip()
    nombre = str(cells[2]).strip()
    costo_proy = _num(cells[3])
    pim = _num(cells[4])
    devengado = _num(cells[5])
    avance_pct = (devengado / pim * 100.0) if pim > 0 else 0.0
    return {
        "id": row.get("id"),
        _COD_DGPP: c_dgpp,
        _COD_SNIP: c_snip,
        _NOM_PROYECTO: nombre,
        _COSTO_PROYECTO: costo_proy,
        "pim": pim,
        "devengado": devengado,
        "avance_porcentaje": round(avance_pct, 2),
        "devengado_acumulado": _num(cells[7]) if len(cells) > 7 else 0.0,
        _SECTOR: detectar_sector(nombre),
        "ficha_url": _ficha_url(c_dgpp, c_snip),
    }


def _fila_banco(row: Dict[str, Any]) -> Dict[str, Any]:
    cells = row.get("cell", [])
    c_dgpp = str(cells[0]).strip()
    c_snip = str(cells[1]).strip()
    nombre = str(cells[2]).strip()
    return {
        "id": row.get("id"),
        _COD_DGPP: c_dgpp,
        _COD_SNIP: c_snip,
        _NOM_PROYECTO: nombre,
        "fecha_registro": str(cells[3]).strip(),
        "fecha_viabilidad": str(cells[4]).strip(),
        "monto_viabilidad": _num(cells[5]),
        _COSTO_PROYECTO: _num(cells[6]),
        _SECTOR: detectar_sector(nombre),
        "ficha_url": _ficha_url(c_dgpp, c_snip),
    }


def _fila_detalle(row: Dict[str, Any]) -> Dict[str, Any]:
    cells = row.get("cell", [])
    return {
        _PERIODO: str(cells[0]),
        "pim": _num(cells[1]),
        "trim1": _num(cells[2]),
        "trim2": _num(cells[3]),
        "trim3": _num(cells[4]),
        "trim4": _num(cells[5]),
        "total_devengado": _num(cells[6]),
        "avance": _num(cells[7]),
    }


def _parsear_filas(raw_rows, minimo_celdas: int, convertir) -> list:
    """Convierte las filas de la grilla; las incompletas o con datos inválidos se omiten."""
    filas = []
    for row in raw_rows:
        if len(row.get("cell", [])) < minimo_celdas:
            continue
        try:
            filas.append(convertir(row))
        except Exception:
            continue
    return filas


class MEFProyService:
    def __init__(self):
        self.ssl_ctx = ssl.create_default_context()
        self.ssl_ctx.check_hostname = False
        self.ssl_ctx.verify_mode = ssl.CERT_NONE

    def _get_headers(self, id_ue: str = "818", periodo: str = "2026") -> Dict[str, str]:
        return {
            "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/122.0.0.0 Safari/537.36",
            "Content-Type": "application/json; charset=utf-8",
            "Accept": "application/json, text/javascript, */*; q=0.01",
            "X-Requested-With": "XMLHttpRequest",
            "Origin": "https://ofi5.mef.gob.pe",
            "Referer": f"{MEF_BASE_URL}?tipo=2&IdUE={id_ue}&IdUEBase={id_ue}&periodoBase={periodo}",
        }

    def _call_webmethod_powershell(
        self, method: str, payload: Dict[str, Any], id_ue: str = "818", periodo: str = "2026"
    ) -> Dict[str, Any]:
        """Ejecuta la petición usando PowerShell Ghost Mode para saltar proxies y restricciones de red (como en Boletín/SEACE)."""
        url = f"{MEF_BASE_URL}/{method}"
        payload_str = json.dumps(payload)
        payload_b64 = base64.b64encode(payload_str.encode(_UTF_8)).decode("ascii")
        referer = f"{MEF_BASE_URL}?tipo=2&IdUE={id_ue}&IdUEBase={id_ue}&periodoBase={periodo}"

        ps_lines = [
            "[Console]::OutputEncoding = [System.Text.Encoding]::UTF8",
            "[System.Net.ServicePointManager]::SecurityProtocol = [System.Net.SecurityProtocolType]::Tls12 -bor [System.Net.SecurityProtocolType]::Tls13",
            "[System.Net.ServicePointManager]::ServerCertificateValidationCallback = {$true}",
            "$wc = New-Object System.Net.WebClient",
            "$wc.Proxy = [System.Net.WebRequest]::GetSystemWebProxy()",
            "$wc.Proxy.Credentials = [System.Net.CredentialCache]::DefaultNetworkCredentials",
            "$wc.Headers.Add('Content-Type', 'application/json; charset=utf-8')",
            "$wc.Headers.Add('Accept', 'application/json, text/javascript, */*; q=0.01')",
            "$wc.Headers.Add('X-Requested-With', 'XMLHttpRequest')",
            "$wc.Headers.Add('User-Agent', 'Mozilla/5.0 (Windows NT 10.0; Win64; x64)')",
            f"$wc.Headers.Add('Referer', '{referer}')",
            f"$body = [System.Text.Encoding]::UTF8.GetString([System.Convert]::FromBase64String('{payload_b64}'))",
            f"$resp = $wc.UploadString('{url}', 'POST', $body)",
            "[Console]::Out.Write($resp)",
        ]
        ps_script = "\n".join(ps_lines)
        encoded_cmd = base64.b64encode(ps_script.encode("utf-16le")).decode("ascii")

        system_root = os.environ.get("SystemRoot", "C:\\Windows")
        ps_exe = os.path.join(system_root, "System32", "WindowsPowerShell", "v1.0", "powershell.exe")
        if not os.path.exists(ps_exe):
            ps_exe = "powershell"

        try:
            proc = subprocess.run(
                [ps_exe, "-NoProfile", "-ExecutionPolicy", "Bypass", "-EncodedCommand", encoded_cmd],
                capture_output=True,
                text=True,
                encoding=_UTF_8,
                errors=_IGNORE,
                timeout=25,
                creationflags=0x08000000,
            )
            if proc.returncode == 0 and proc.stdout.strip():
                return json.loads(proc.stdout.strip())
            raise Exception(f"PowerShell WebClient falló: {proc.stderr.strip() or f'RC={proc.returncode}'}")
        except subprocess.TimeoutExpired as exc:
            raise TimeoutError(
                "El servidor del MEF (ofi5.mef.gob.pe) excedió el tiempo máximo de espera (Timeout de 25s en PowerShell)."
            ) from exc

    def _call_webmethod_curl(
        self, method: str, payload: Dict[str, Any], id_ue: str = "818", periodo: str = "2026"
    ) -> Dict[str, Any]:
        """Ejecuta la petición usando curl.exe nativo de Windows (HTTP/2 + Schannel) para velocidad óptima y bypass de WAF F5 del MEF."""
        url = f"{MEF_BASE_URL}/{method}"
        payload_str = json.dumps(payload)
        referer = f"{MEF_BASE_URL}?tipo=2&IdUE={id_ue}&IdUEBase={id_ue}&periodoBase={periodo}"

        system_root = os.environ.get("SystemRoot", "C:\\Windows")
        curl_exe = os.path.join(system_root, "System32", "curl.exe")
        if not os.path.exists(curl_exe):
            curl_exe = "curl"

        cmd = [
            curl_exe,
            "-4",
            "--connect-timeout",
            "8",
            "--max-time",
            "20",
            "--ssl-no-revoke",
            "-s",
            "-k",
            "-X",
            "POST",
            url,
            "-H",
            "Content-Type: application/json; charset=utf-8",
            "-H",
            "Accept: application/json, text/javascript, */*; q=0.01",
            "-H",
            "X-Requested-With: XMLHttpRequest",
            "-H",
            "User-Agent: Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/122.0.0.0 Safari/537.36",
            "-H",
            f"Referer: {referer}",
            "-d",
            payload_str,
        ]

        proc = subprocess.run(
            cmd, capture_output=True, text=True, encoding=_UTF_8, errors=_IGNORE, timeout=25, creationflags=0x08000000
        )
        if proc.returncode == 0 and proc.stdout.strip():
            raw = proc.stdout.strip()
            if raw.startswith("{") and raw.endswith("}"):
                return json.loads(raw)
            raise Exception(f"Respuesta de curl no es JSON válido: {raw[:150]}")
        else:
            raise Exception(f"curl.exe falló: {proc.stderr.strip() or f'RC={proc.returncode}'}")

    def _call_webmethod(
        self, method: str, payload: Dict[str, Any], id_ue: str = "818", periodo: str = "2026"
    ) -> Dict[str, Any]:
        """Consulta WebMethod con triple motor y reintentos automáticos contra congestionamiento del MEF."""
        import time

        for intento in range(1, 3):
            # 1. Motor primario ultra-rápido: curl.exe nativo (-4 IPv4)
            try:
                return self._call_webmethod_curl(method, payload, id_ue=id_ue, periodo=periodo)
            except Exception as err_curl:
                logger.warning(
                    f"[Intento {intento}] Motor curl nativo falló ({err_curl}). Intentando conexión directa urllib..."
                )

            # 2. Motor secundario: urllib directa
            try:
                url = f"{MEF_BASE_URL}/{method}"
                data_bytes = json.dumps(payload).encode(_UTF_8)
                headers = self._get_headers(id_ue=id_ue, periodo=periodo)

                req = urllib.request.Request(url, data=data_bytes, headers=headers)
                with urllib.request.urlopen(req, context=self.ssl_ctx, timeout=12) as resp:
                    raw = resp.read().decode("utf-8-sig", errors=_IGNORE)
                    return json.loads(raw)
            except Exception as err_urllib:
                logger.warning(
                    f"[Intento {intento}] Conexión urllib falló ({err_urllib}). Conmutando a modo PowerShell..."
                )

            # 3. Motor terciario: PowerShell Ghost Mode
            try:
                return self._call_webmethod_powershell(method, payload, id_ue=id_ue, periodo=periodo)
            except Exception as err_ps:
                logger.warning(f"[Intento {intento}] Motor PowerShell falló: {err_ps}")

            time.sleep(1)

        raise TimeoutError(
            "La plataforma pública del MEF (ofi5.mef.gob.pe) se encuentra temporalmente lenta o saturada. Por favor vuelve a hacer clic en 'Consultar'."
        )

    def obtener_presets_ue(self) -> List[Dict[str, str]]:
        return UNIDADES_EJECUTORAS_PRESET

    def obtener_periodos(self, p_periodo: str = "Ejecucion") -> List[str]:
        try:
            res = self._call_webmethod("GetPeriodo", {_PPERIODO: p_periodo})
            items = res.get("d", [])
            periodos = [item.get("Valor") for item in items if item.get("Valor")]
            periodos.sort(reverse=True)
            return periodos if periodos else ["2026", "2025", "2024", "2023", "2022", "2021", "2020"]
        except Exception:
            return ["2026", "2025", "2024", "2023", "2022", "2021", "2020"]

    def obtener_ejecucion(
        self,
        id_ue: str = "818",
        periodo: str = "2026",
        page_size: int = 50,
        page_number: int = 1,
        sort_column: str = "codDGPP",
        sort_order: str = "asc",
        nom_proyecto: str = "",
        cod_dgpp: str = "",
        cod_snip: str = "",
    ) -> Dict[str, Any]:
        """Consulta la ejecución presupuestal de proyectos en el MEF."""
        payload = {
            _PPAGESIZE: page_size,
            _PPAGENUMBER: page_number,
            _PSORTCOLUMN: sort_column,
            _PSORTORDER: sort_order,
            _PPERIODO: periodo,
            _PCODEJECUTORA: id_ue,
            "pCodDGPP": cod_dgpp,
            "pCodSNIP": cod_snip,
            "pNomProyecto": nom_proyecto,
        }

        res = self._call_webmethod("GetEjecucion", payload, id_ue=id_ue, periodo=periodo)
        data_grid = res.get("d", {})

        raw_rows = data_grid.get("rows", [])
        total_records = data_grid.get("records", 0)
        total_pages = data_grid.get("total", 0)
        current_page = data_grid.get("page", page_number)

        parsed_rows = _parsear_filas(raw_rows, 8, _fila_ejecucion)
        total_costo_p = sum(f[_COSTO_PROYECTO] for f in parsed_rows)
        total_pim_p = sum(f["pim"] for f in parsed_rows)
        total_devengado_p = sum(f["devengado"] for f in parsed_rows)

        avance_global = (total_devengado_p / total_pim_p * 100.0) if total_pim_p > 0 else 0.0

        return {
            "status": "success",
            "id_ue": id_ue,
            _PERIODO: periodo,
            "page": current_page,
            "total_pages": total_pages,
            "total_records": total_records,
            "page_records": len(parsed_rows),
            "totales_pagina": {
                "pim_total": round(total_pim_p, 2),
                "devengado_total": round(total_devengado_p, 2),
                "avance_global": round(avance_global, 2),
                "costo_proyectos_total": round(total_costo_p, 2),
            },
            "proyectos": parsed_rows,
        }

    def obtener_banco_proyectos(
        self,
        id_ue: str = "818",
        periodo: str = "2026",
        page_size: int = 50,
        page_number: int = 1,
        sort_column: str = "codDGPP",
        sort_order: str = "asc",
    ) -> Dict[str, Any]:
        """Consulta el Banco de Proyectos (DGPI) con fechas de registro y viabilidad."""
        payload = {
            _PPAGESIZE: page_size,
            _PPAGENUMBER: page_number,
            _PSORTCOLUMN: sort_column,
            _PSORTORDER: sort_order,
            _PPERIODO: periodo,
            _PCODEJECUTORA: id_ue,
        }

        res = self._call_webmethod("GetProyectos", payload, id_ue=id_ue, periodo=periodo)
        data_grid = res.get("d", {})
        raw_rows = data_grid.get("rows", [])
        total_records = data_grid.get("records", 0)
        total_pages = data_grid.get("total", 0)

        proyectos = _parsear_filas(raw_rows, 7, _fila_banco)

        return {
            "status": "success",
            "id_ue": id_ue,
            _PERIODO: periodo,
            "total_records": total_records,
            "total_pages": total_pages,
            "proyectos": proyectos,
        }

    def obtener_ranking(self, id_ue: str = "818", periodo: str = "2026") -> List[Dict[str, Any]]:
        try:
            res = self._call_webmethod(
                "GetRankingProyecto", {_PPERIODO: periodo, _PCODEJECUTORA: id_ue}, id_ue=id_ue, periodo=periodo
            )
            items = res.get("d", [])
            ranking = []
            for it in items:
                nombre = it.get(_NOMBRE, "")
                ranking.append(
                    {
                        _COD_DGPP: it.get("dgpp", ""),
                        _COD_SNIP: it.get("snip", ""),
                        _NOMBRE: nombre,
                        "pim": float(it.get("pim", 0)),
                        _SECTOR: detectar_sector(nombre),
                    }
                )
            return ranking
        except Exception:
            return []

    def obtener_detalle_proyecto(self, id_ue: str, cod_proyecto: str) -> List[Dict[str, Any]]:
        payload = {
            _PPAGESIZE: 50,
            _PPAGENUMBER: 1,
            _PSORTCOLUMN: _PERIODO,
            _PSORTORDER: "desc",
            _PCODEJECUTORA: id_ue,
            "pCodProyecto": cod_proyecto,
        }

        res = self._call_webmethod("GetDetalle", payload, id_ue=id_ue)
        data_grid = res.get("d", {})
        raw_rows = data_grid.get("rows", [])

        detalles = _parsear_filas(raw_rows, 8, _fila_detalle)

        return detalles

    def generar_excel_xlsx(
        self, proyectos: List[Dict[str, Any]], id_ue: str, periodo: str, tipo_vista: str = _EJECUCION
    ) -> bytes:
        """Genera un archivo XLSX ejecutivo de alto impacto visual con cabecera agrupada, fórmulas y formato financiero."""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = f"Proyectos MEF {periodo}"
        ws.views.sheetView[0].showGridLines = True

        num_cols = 8 if tipo_vista == _EJECUCION else 7
        _xlsx_encabezado(ws, id_ue, periodo, tipo_vista, num_cols)
        if tipo_vista == _EJECUCION:
            _xlsx_cabecera_ejecucion(ws, periodo)
        else:
            _xlsx_cabecera_banco(ws)
        ws.row_dimensions[_XLSX_R5].height = 22
        ws.row_dimensions[_XLSX_R6].height = 22

        curr_row = 7
        for idx, p in enumerate(proyectos):
            _xlsx_fila_datos(ws, curr_row, idx, p, tipo_vista)
            curr_row += 1

        if tipo_vista == _EJECUCION and curr_row > 7:
            _xlsx_fila_totales(ws, curr_row)

        for col_l, width in _XLSX_COL_WIDTHS.items():
            if ord(col_l) - ord("A") + 1 <= num_cols:
                ws.column_dimensions[col_l].width = width

        buffer = io.BytesIO()
        wb.save(buffer)
        return buffer.getvalue()


# ── Estilos y bloques del Excel ejecutivo ─────────────────────────────────────
_XLSX_R5 = 5
_XLSX_R6 = 6

# Paleta armónica con la grilla
_C_HDR_BG = "E4F1FC"  # Azul pastel suave
_C_HDR_GROUP_BG = "D4ECFC"  # Azul pastel para grupo año
_C_HDR_TEXT = "033F6C"  # Azul texto institucional
_C_HDR_BORDER = "BEDBFF"  # Borde azul claro
_C_TOTAL_BG = "EBF4FC"  # Fondo totales
_C_ZEBRA = "F8FAFC"  # Fila cebra
_FMT_SOLES = '"S/" #,##0.00'

_HEADER_FONT = Font(name=_CALIBRI, size=10, bold=True, color=_C_HDR_TEXT)
_SUB_HEADER_FONT = Font(name=_CALIBRI, size=9.5, bold=True, color=_C_HDR_TEXT)
_GROUP_FONT = Font(name=_CALIBRI, size=10.5, bold=True, color=_C_HDR_TEXT)
_TITLE_FONT = Font(name=_CALIBRI, size=13, bold=True, color="FFFFFF")
_SUBTITLE_FONT = Font(name=_CALIBRI, size=11, bold=True, color="014F88")
_META_FONT = Font(name=_CALIBRI, size=9, bold=False, color="475569")
_BOLD_FONT = Font(name=_CALIBRI, size=9.5, bold=True, color="0F172A")
_REGULAR_FONT = Font(name=_CALIBRI, size=9.5, color="1E293B")
_CODE_FONT = Font(name="Consolas", size=9.5, bold=True, color=_K_0163AC)


def _side(style, color):
    return Side(style=style, color=color)


_HEADER_BORDER = Border(
    left=_side("thin", _C_HDR_BORDER),
    right=_side("thin", _C_HDR_BORDER),
    top=_side("thin", _C_HDR_BORDER),
    bottom=_side("thin", _C_HDR_BORDER),
)
_THIN_BORDER = Border(
    left=_side("thin", _E2E8F0),
    right=_side("thin", _E2E8F0),
    top=_side("thin", _E2E8F0),
    bottom=_side("thin", _E2E8F0),
)
_TOTAL_TOP_BORDER = Border(
    left=_side("thin", _C_HDR_BORDER),
    right=_side("thin", _C_HDR_BORDER),
    top=_side("medium", _K_0163AC),
    bottom=_side("double", _K_0163AC),
)

_XLSX_COL_WIDTHS = {
    "A": 16,  # CÓDIGO ÚNICO
    "B": 15,  # CÓDIGO SNIP
    "C": 55,  # PROYECTO DE INVERSIÓN (espacio amplio)
    "D": 22,  # COSTO DEL PROYECTO
    "E": 20,  # PIM
    "F": 20,  # DEVENGADO
    "G": 15,  # AVANCE %
    "H": 26,  # DEVENGADO ACUMULADO
}


def _solid(color):
    return PatternFill(start_color=color, fill_type="solid")


def _center(wrap=False):
    if wrap:
        return Alignment(horizontal=_CENTER, vertical=_CENTER, wrap_text=True)
    return Alignment(horizontal=_CENTER, vertical=_CENTER)


def _xlsx_banner_row(ws, row, last_col, value, font, height, fill=None):
    ws.merge_cells(f"A{row}:{last_col}{row}")
    cell = ws[f"A{row}"]
    cell.value = value
    cell.font = font
    if fill:
        cell.fill = fill
    cell.alignment = _center()
    ws.row_dimensions[row].height = height


def _xlsx_encabezado(ws, id_ue, periodo, tipo_vista, num_cols):
    ue_info = next((u for u in UNIDADES_EJECUTORAS_PRESET if u["id"] == id_ue), None)
    inst_nombre = ue_info[_NOMBRE] if ue_info else f"UNIDAD EJECUTORA N° {id_ue}"
    dep_nombre = ue_info[_DEPARTAMENTO] if ue_info else ""
    last_col = get_column_letter(num_cols)

    banner_fill = PatternFill(start_color=_K_0163AC, end_color=_K_0163AC, fill_type="solid")
    _xlsx_banner_row(
        ws, 1, last_col, "SISTEMA DE CONSULTA DE INVERSIÓN PÚBLICA (MEF - SIAF / PTE)", _TITLE_FONT, 28, banner_fill
    )
    inst = f"GOBIERNO REGIONAL / ENTIDAD: {inst_nombre}" + (f" ({dep_nombre.upper()})" if dep_nombre else "")
    _xlsx_banner_row(ws, 2, last_col, inst, _SUBTITLE_FONT, 20)
    meta = (
        f"UNIDAD EJECUTORA: {id_ue} | AÑO FISCAL: {periodo} | VISTA: {tipo_vista.upper()} | "
        f"FECHA DE EMISIÓN: {datetime.now().strftime('%d/%m/%Y %H:%M')}"
    )
    _xlsx_banner_row(ws, 3, last_col, meta, _META_FONT, 18)


def _xlsx_header_cell(ws, ref, value, fill_color, font):
    cell = ws[ref]
    cell.value = value
    cell.fill = _solid(fill_color)
    cell.font = font
    cell.alignment = _center(wrap=True)
    return cell


def _xlsx_cabecera_ejecucion(ws, periodo):
    r5, r6 = _XLSX_R5, _XLSX_R6
    headers_r5 = {
        "A": "CÓDIGO\nÚNICO",
        "B": "CÓDIGO\nSNIP",
        "C": "PROYECTO",
        "D": "COSTO DEL\nPROYECTO",
        "H": f"DEVENGADO\nACUMULADO\nAL {periodo}",
    }
    for col_l, title in headers_r5.items():
        ws.merge_cells(f"{col_l}{r5}:{col_l}{r6}")
        _xlsx_header_cell(ws, f"{col_l}{r5}", title, _C_HDR_BG, _HEADER_FONT)

    # Grupo del año: E5:G5
    ws.merge_cells(f"E{r5}:G{r5}")
    cell_group = ws[f"E{r5}"]
    cell_group.value = f"{periodo}"
    cell_group.fill = _solid(_C_HDR_GROUP_BG)
    cell_group.font = _GROUP_FONT
    cell_group.alignment = _center()

    for col_l, title in {"E": "PIM", "F": "DEVENGADO", "G": "AVANCE\n(%)"}.items():
        _xlsx_header_cell(ws, f"{col_l}{r6}", title, _C_HDR_BG, _SUB_HEADER_FONT)

    for r in [r5, r6]:
        for c in range(1, 9):
            ws.cell(row=r, column=c).border = _HEADER_BORDER


def _xlsx_cabecera_banco(ws):
    headers_banco = [
        "CÓDIGO ÚNICO",
        "CÓDIGO SNIP",
        "PROYECTO",
        "FECHA REGISTRO",
        "FECHA VIABILIDAD",
        "MONTO VIABILIDAD",
        "COSTO TOTAL",
    ]
    for col_idx, h in enumerate(headers_banco, 1):
        col_l = get_column_letter(col_idx)
        ws.merge_cells(f"{col_l}{_XLSX_R5}:{col_l}{_XLSX_R6}")
        _xlsx_header_cell(ws, f"{col_l}{_XLSX_R5}", h, _C_HDR_BG, _HEADER_FONT)
        for r in [_XLSX_R5, _XLSX_R6]:
            ws.cell(row=r, column=col_idx).border = _HEADER_BORDER


def _xlsx_datos_proyecto(p, tipo_vista):
    codigo = p.get(_COD_DGPP, "") or p.get(_COD_SNIP, "")
    if tipo_vista == "banco":
        return [
            codigo,
            p.get(_COD_SNIP, ""),
            p.get(_NOM_PROYECTO, ""),
            p.get("fecha_registro", ""),
            p.get("fecha_viabilidad", ""),
            float(p.get("monto_viabilidad", 0) or 0),
            float(p.get(_COSTO_PROYECTO, 0) or 0),
        ]
    pim_val = float(p.get("pim", 0) or 0)
    dev_val = float(p.get("devengado", 0) or 0)
    av_val = (dev_val / pim_val) if pim_val > 0 else 0.0
    return [
        codigo,
        p.get(_COD_SNIP, ""),
        p.get(_NOM_PROYECTO, ""),
        float(p.get(_COSTO_PROYECTO, 0) or 0),
        pim_val,
        dev_val,
        av_val,
        float(p.get("devengado_acumulado", 0) or 0),
    ]


def _xlsx_estilo_valor(cell, col_idx, tipo_vista):
    """Fuente, alineación y formato numérico de una celda de datos."""
    if col_idx == 1:
        cell.font = _CODE_FONT
        cell.alignment = _center()
        return
    cell.font = _REGULAR_FONT
    if col_idx == 2:
        cell.alignment = _center()
    elif col_idx == 3:
        cell.alignment = Alignment(horizontal="left", vertical=_CENTER, wrap_text=True)
    elif tipo_vista == _EJECUCION:
        if col_idx in [4, 5, 6, 8]:
            cell.number_format = _FMT_SOLES
            cell.alignment = Alignment(horizontal=_RIGHT, vertical=_CENTER)
        elif col_idx == 7:
            cell.number_format = "0.0%"
            cell.alignment = _center()
    elif col_idx in [6, 7]:
        cell.number_format = _FMT_SOLES
        cell.alignment = Alignment(horizontal=_RIGHT, vertical=_CENTER)
    elif col_idx in [4, 5]:
        cell.alignment = _center()


def _xlsx_fila_datos(ws, curr_row, idx, p, tipo_vista):
    row_fill = _solid(_C_ZEBRA if idx % 2 == 1 else "FFFFFF")
    for col_idx, val in enumerate(_xlsx_datos_proyecto(p, tipo_vista), 1):
        cell = ws.cell(row=curr_row, column=col_idx, value=val)
        cell.fill = row_fill
        cell.border = _THIN_BORDER
        _xlsx_estilo_valor(cell, col_idx, tipo_vista)
    ws.row_dimensions[curr_row].height = 20


def _xlsx_fila_totales(ws, curr_row):
    ws.merge_cells(start_row=curr_row, start_column=1, end_row=curr_row, end_column=3)
    tot_cell = ws.cell(row=curr_row, column=1, value="TOTAL GENERAL:")
    tot_cell.font = _BOLD_FONT
    tot_cell.alignment = Alignment(horizontal=_RIGHT, vertical=_CENTER)

    for c in range(1, 4):
        cell = ws.cell(row=curr_row, column=c)
        cell.border = _TOTAL_TOP_BORDER
        cell.fill = _solid(_C_TOTAL_BG)

    for col_idx in [4, 5, 6, 8]:
        col_letter = get_column_letter(col_idx)
        cell = ws.cell(row=curr_row, column=col_idx)
        cell.value = f"=SUM({col_letter}7:{col_letter}{curr_row - 1})"
        cell.font = _BOLD_FONT
        cell.number_format = _FMT_SOLES
        cell.border = _TOTAL_TOP_BORDER
        cell.fill = _solid(_C_TOTAL_BG)
        cell.alignment = Alignment(horizontal=_RIGHT, vertical=_CENTER)

    cell_av = ws.cell(row=curr_row, column=7)
    cell_av.value = f"=IF(E{curr_row}>0, F{curr_row}/E{curr_row}, 0)"
    cell_av.font = _BOLD_FONT
    cell_av.number_format = "0.0%"
    cell_av.border = _TOTAL_TOP_BORDER
    cell_av.fill = _solid(_C_TOTAL_BG)
    cell_av.alignment = _center()

    ws.row_dimensions[curr_row].height = 24
