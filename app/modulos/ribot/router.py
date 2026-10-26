import os
import re
import uuid
import time
import shutil
import zipfile
import json
from decimal import Decimal
from typing import List
from fastapi import APIRouter, Request, Depends, UploadFile, File, Form
from fastapi.concurrency import run_in_threadpool
from fastapi.responses import JSONResponse, FileResponse
from fastapi.templating import Jinja2Templates
import pdfplumber
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

from app.modulos.auth.dependencies import login_required
from app.core.config import settings, sicode_dir
from app.modulos.ribot.padron import PadronError, PadronUnavailableError, formato_domicilio, get_padron
from app.modulos.ribot.resolutivo import adapt_resolutivo
from app.modulos.ribot.texto_libre import (
    MARKER as TEXTO_LIBRE_MARCADOR,
    MAX_LENGTH as TEXTO_MAX,
    apply_texto_libre as aplicar_texto_libre,
    normalize_lines as normalizar_lineas,
)

# Textos repetidos
_AVISOS = "avisos"
_AVISOS_EXTRACCION = "avisos_extraccion"
_CIUDAD = "ciudad"
_CODIGO_TRIBUTO = "codigo_tributo"
_CODIGO_TRIBUTO_ASOC = "codigo_tributo_asoc"
_COMPLETAR = "[COMPLETAR]"
_COUNT = "count"
_DOCX = ".docx"
_DOMICILIO = "domicilio"
_DOWNLOAD_URL = "download_url"
_ERROR = "error"
_FECHA_DECLARACION = "fecha_declaracion"
_FECHA_EMISION_OP = "fecha_emision_op"
_FECHA_EXPEDIENTE = "fecha_expediente"
_FECHA_INTERES = "fecha_interes"
_FECHA_PAGO = "fecha_pago"
_FECHA_RESOLUCION = "fecha_resolucion"
_FILENAME = "filename"
_FOLDERS = "folders"
_FORMULARIO_PDT = "formulario_pdt"
_FUNDADO = "fundado"
_FUNDADO_2 = "FUNDADO"
_HUANCAYO = "Huancayo"
_IMPUESTO_GENERAL_A_LAS_VENTA = "Impuesto General a las Ventas"
_INTENDENCIA = "intendencia"
_INTERES_MODIFICADO = "interes_modificado"
_INTERES_MULTA = "interes_multa"
_INTERES_MULTA_MOD = "interes_multa_mod"
_INTERES_ORIGINAL = "interes_original"
_MESSAGE = "message"
_MULTA = "multa"
_MULTA_MOD = "multa_mod"
_NOMBRE_TRIBUTO = "nombre_tributo"
_NRO_BOLETA_PAGO = "nro_boleta_pago"
_NRO_EXPEDIENTE = "nro_expediente"
_NRO_ORDEN_DECLARACION = "nro_orden_declaracion"
_NRO_ORDEN_PAGO = "nro_orden_pago"
_NRO_RESOLUCION = "nro_resolucion"
_PADRON_CAIDO = "padron_caido"
_PERIODO = "periodo"
_PHYSICAL_PATH = "physical_path"
_RAZON_SOCIAL = "razon_social"
_SAVED_LOCALLY = "saved_locally"
_STATUS = "status"
_SUCCESS = "success"
_TIPO_VALOR = "tipo_valor"
_TOTAL_MODIFICADO = "total_modificado"
_TOTAL_MULTA = "total_multa"
_TOTAL_MULTA_MOD = "total_multa_mod"
_TOTAL_ORIGINAL = "total_original"
_TRIBUTO_MODIFICADO = "tributo_modificado"
_TRIBUTO_ORIGINAL = "tributo_original"

router = APIRouter(prefix="/ribot", tags=["ribot"], dependencies=[Depends(login_required)])

templates_dir = settings.get_template_path("app/modulos/ribot/templates")
templates_base = settings.get_template_path("app/modulos/inicio/templates")
templates = Jinja2Templates(directory=[templates_dir, templates_base])

# Directorio temporal de RIbot
TEMP_DIR = sicode_dir("DATA", "RIBOT", "TEMP")

# Directorio central de documentos locales (Estilo Bpsearch)
RIBOT_BASE_DIR = sicode_dir("DATA", "RIBOT", "DOCS")


# Directorio de Plantillas Word de RIbot
def obtener_directorio_plantillas() -> str:
    candidato = r"D:\SICODE\APPS\RIBOT"
    if os.path.exists(candidato):
        return candidato
    local = os.path.join(os.path.dirname(__file__), "plantillas")
    os.makedirs(local, exist_ok=True)
    return local


RIBOT_PLANTILLAS_DIR = obtener_directorio_plantillas()


# ── Sanitización y Limpieza de Archivos ──────────────────────────────────────
def sanitizar_nombre_archivo(nombre: str) -> str:
    r"""Remueve caracteres inválidos en Windows (\ / : * ? " < > |) para evitar crashes OSError."""
    if not nombre:
        return "DOCUMENTO"
    limpio = re.sub(r'[\\/:*?"<>|\r\n\t]', "_", str(nombre))
    limpio = re.sub(r"_+", "_", limpio).strip(" ._")
    return limpio or "DOCUMENTO"


def limpiar_archivos_temporales(max_horas=6):
    """Elimina carpetas temporales antiguas en TEMP_DIR para evitar fugas de almacenamiento."""
    try:
        ahora = time.time()
        limite = ahora - (max_horas * 3600)
        for item in os.listdir(TEMP_DIR):
            item_path = os.path.join(TEMP_DIR, item)
            if os.path.isdir(item_path):
                st_time = os.path.getmtime(item_path)
                if st_time < limite:
                    shutil.rmtree(item_path, ignore_errors=True)
    except Exception as e:
        print(f"Error en rutina de limpieza de temporales RIbot: {e}")


# ── Configuración de Fechas SUNAT ──────────────────────────────────────────────
MESES = {
    "01": "enero",
    "02": "febrero",
    "03": "marzo",
    "04": "abril",
    "05": "mayo",
    "06": "junio",
    "07": "julio",
    "08": "agosto",
    "09": "septiembre",
    "10": "octubre",
    "11": "noviembre",
    "12": "diciembre",
}


def fecha_larga(fecha_str):
    """Convierte '09/03/2026' → '09 de marzo de 2026' de forma segura."""
    if not fecha_str or not isinstance(fecha_str, str):
        return _COMPLETAR
    try:
        partes = fecha_str.strip().split("/")
        if len(partes) == 3:
            d, m, a = partes
            mes_nombre = MESES.get(m.zfill(2), m)
            return f"{d} de {mes_nombre} de {a}"
    except Exception:
        pass
    return fecha_str


# ── Auxiliares para Generación de Word (python-docx) ──────────────────────────
def set_cell_bg(cell, hex_color):
    tc = cell._tc
    tc_pr = tc.get_or_add_tc_pr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tc_pr.append(shd)


def set_cell_borders(cell):
    tc = cell._tc
    tc_pr = tc.get_or_add_tc_pr()
    borders = OxmlElement("w:tcBorders")
    for side in ["top", "left", "bottom", "right"]:
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "000000")
        borders.append(el)
    tc_pr.append(borders)


def agregar_tabla(doc, encabezados, fila_datos, col_widths_cm):
    table = doc.add_table(rows=2, cols=len(encabezados))
    table.style = "Table Grid"

    # Fila encabezado
    for i, texto in enumerate(encabezados):
        cell = table.rows[0].cells[i]
        cell.width = Cm(col_widths_cm[i])
        set_cell_bg(cell, "D9D9D9")
        set_cell_borders(cell)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(texto)
        run.bold = True
        run.font.size = Pt(9)

    # Fila datos
    for i, texto in enumerate(fila_datos):
        cell = table.rows[1].cells[i]
        cell.width = Cm(col_widths_cm[i])
        set_cell_borders(cell)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(str(texto))
        if i == 0:
            run.bold = True
        run.font.size = Pt(9)

    return table


def parrafo_justificado(doc, partes, sangria_cm=0):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if sangria_cm:
        p.paragraph_format.first_line_indent = Cm(sangria_cm)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.space_before = Pt(0)
    for texto, bold in partes:
        run = p.add_run(texto)
        run.bold = bold
        run.font.size = Pt(10)
    return p


NOMBRES_TRIBUTO = {
    "010101": _IMPUESTO_GENERAL_A_LAS_VENTA,
    "011001": "Impuesto a la Renta de Tercera Categoría",
    "053100": "SNP Ley 19990",
}

MAX_PAGINAS_PDF = 30  # una Ficha del Valor tiene pocas páginas; evita que un PDF enorme acapare el servidor


# ── Extractor Ficha del Valor PDF ───────────────────────────────────────────
def _ficha_seccion_1(texto, datos):
    # 1. Intendencia y Ciudad
    m = re.search(r"^(I\.T\.I\.[^\n]+)", texto, re.MULTILINE)
    if m:
        datos[_INTENDENCIA] = m.group(1).split("Ficha del Valor")[0].strip()
    else:
        m_alt = re.search(r"^(.+?)\s+Ficha del Valor", texto, re.MULTILINE)
        if m_alt:
            datos[_INTENDENCIA] = m_alt.group(1).strip()

    intend = datos.get(_INTENDENCIA, "").upper()
    if "JUNIN" in intend or "HUANCAYO" in intend:
        datos[_CIUDAD] = _HUANCAYO
    elif "LIMA" in intend:
        datos[_CIUDAD] = "Lima"
    else:
        datos[_CIUDAD] = _HUANCAYO


_PATRON_RUC = r"(?:10|15|16|17|20)\d{9}"
_CORTE_RAZON_SOCIAL = r"(?:Notificaci|Montos|Periodo|REC|Estado)"
_SECCION_SIGUIENTE = r"(?:Notificaci|Montos|Periodo|REC|Estado|Datos|C[aá]lculo|Fecha|Tributo|Emisi|Etapa)"


def _continuacion_razon_social(texto, m_ruc):
    """La razón social puede continuar en la línea siguiente (ej. '... SR. DE' / 'ATACO SOCIEDAD ANONIMA CERRADA')."""
    lineas = texto[m_ruc.end() :].lstrip("\r\n").splitlines()
    if not lineas:
        return ""
    siguiente = lineas[0].strip()
    if re.fullmatch(r"[A-ZÁÉÍÓÚÑ0-9&.,\-()/ ]{3,}", siguiente) and not re.match(
        _SECCION_SIGUIENTE, siguiente, re.IGNORECASE
    ):
        return " " + siguiente
    return ""


def _razon_social_de_linea_ruc(texto, m_ruc):
    candidata = m_ruc.group(2).strip()
    candidata = re.split(_CORTE_RAZON_SOCIAL, candidata, flags=re.IGNORECASE)[0].strip()
    candidata += _continuacion_razon_social(texto, m_ruc)
    return re.sub(r"\s+", " ", candidata)


def _ruc_suelto(texto, datos):
    """RUC sin etiqueta: la razón social es lo que sigue al RUC en su misma línea."""
    m_ruc2 = re.search(rf"\b({_PATRON_RUC})\b", texto)
    if not m_ruc2:
        return
    datos["ruc"] = m_ruc2.group(1).strip()
    for linea in texto.splitlines():
        if datos["ruc"] in linea:
            partes = linea.split(datos["ruc"])
            if len(partes) > 1 and partes[1].strip():
                datos[_RAZON_SOCIAL] = re.sub(r"\s+", " ", partes[1].strip())
                break


def _ficha_seccion_2(texto, datos):
    # 2. RUC y Razón Social (En Ficha del Valor viene en la misma línea: RUC 20602398871 HICOGROUP S.A.C.)
    m_ruc = re.search(rf"RUC\s*[:\s]*\b({_PATRON_RUC})\b\s*([^\n\r]+)", texto, re.IGNORECASE)
    if m_ruc:
        datos["ruc"] = m_ruc.group(1).strip()
        datos[_RAZON_SOCIAL] = _razon_social_de_linea_ruc(texto, m_ruc)
    else:
        _ruc_suelto(texto, datos)


def _ficha_seccion_3(texto, datos):
    # 3. N° Orden de Pago (13 dígitos)
    m_op = re.search(r"Valor\s+(?:Nro\.?\s*)?(\d{13})", texto, re.IGNORECASE)
    if not m_op:
        m_op = re.search(r"Valor\s+(\d{13})\s+Orden\s+de\s+Pago", texto, re.IGNORECASE)
    if not m_op:
        m_op = re.search(r"Valor Nro\.\s+(\d+)", texto)
    if m_op:
        datos[_NRO_ORDEN_PAGO] = m_op.group(1).strip()

    # Tipo de documento: lo que dice la línea del Valor ("Orden de Pago", "Resolución de Multa", ...)
    m_tipo = re.search(
        r"Valor\s+(?:Nro\.?\s*)?\d{13}\s+([A-Za-zÁÉÍÓÚÜÑáéíóúüñ]+(?:[ \t]+(?:de|del|la|las|los|[A-Za-zÁÉÍÓÚÜÑáéíóúüñ]+))*?)"
        r"(?=[ \t]+(?:Periodo|Tributo|RUC|Estado|Emisi|Etapa|Fecha)\b|[ \t]*\d|[ \t]*$)",
        texto,
        re.IGNORECASE | re.MULTILINE,
    )
    if m_tipo:
        datos[_TIPO_VALOR] = re.sub(r"\s+", " ", m_tipo.group(1)).strip()


def _ficha_seccion_4(texto, datos):
    # 4. Periodo
    m_per = re.search(r"Periodo\s+(\d{4}-\d{2})", texto, re.IGNORECASE)
    if m_per:
        datos[_PERIODO] = m_per.group(1).strip()


def _ficha_seccion_5(texto, datos):
    # 5. Fecha emisión OP
    m_em = re.search(r"Emisi[oó]n\s+(\d{2}/\d{2}/\d{4})", texto, re.IGNORECASE)
    if m_em:
        datos[_FECHA_EMISION_OP] = m_em.group(1).strip()
        datos[_FECHA_DECLARACION] = m_em.group(1).strip()
        datos[_FECHA_INTERES] = m_em.group(1).strip()

    m_fint = re.search(r"Inicio\s+Inter[eé]s\s+(\d{2}/\d{2}/\d{4})", texto, re.IGNORECASE)
    if m_fint:
        datos[_FECHA_INTERES] = m_fint.group(1).strip()

    # Tributo Asociado (solo en Resoluciones de Multa): "Tributo Asociado 053100 PENREG"
    m_asoc = re.search(r"Tributo\s+Asociado\s+(\d{6})", texto, re.IGNORECASE)
    if m_asoc:
        datos[_CODIGO_TRIBUTO_ASOC] = m_asoc.group(1).strip()


def _ficha_seccion_6(texto, datos):
    # 6. Código de Tributo
    m_cod = re.search(r"Tributo\s+(\d{6})", texto, re.IGNORECASE)
    if not m_cod:
        m_cod = re.search(r"\b(\d{6})\s+[A-Z0-9]{3,}", texto)
    if m_cod:
        cod = m_cod.group(1).strip()
        datos[_CODIGO_TRIBUTO] = cod
        datos[_NOMBRE_TRIBUTO] = {
            "010101": _IMPUESTO_GENERAL_A_LAS_VENTA,
            "011001": "Impuesto a la Renta de Tercera Categoría",
            "053100": "SNP Ley 19990",
        }.get(cod, f"Tributo {cod}")
    else:
        datos[_CODIGO_TRIBUTO] = "010101"
        datos[_NOMBRE_TRIBUTO] = _IMPUESTO_GENERAL_A_LAS_VENTA


def _ficha_seccion_7(texto, datos):
    # 7. Montos Originales
    m_montos = re.search(
        r"C[aá]lculo\s+Saldo\s+Total.*?\nFecha\s+(\d{2}/\d{2}/\d{4}).*?\nTributo\s+([\d,]+\.\d{2}).*?\n.*?Inter[eé]s\s+([\d,]+\.\d{2}).*?\nTotal\s+([\d,]+\.\d{2})",
        texto,
        re.IGNORECASE | re.DOTALL,
    )
    if m_montos:
        datos[_TRIBUTO_ORIGINAL] = m_montos.group(2).strip()
        datos[_INTERES_ORIGINAL] = m_montos.group(3).strip()
        datos[_TOTAL_ORIGINAL] = m_montos.group(4).strip()
    else:
        m_op_ev = re.search(
            r"Emisi[oó]n\s+de\s+OP\s+\d+\s+([\d,]+\.\d{2})\s+[\d,]+\.\d{2}\s+([\d,]+\.\d{2})\s+([\d,]+\.\d{2})",
            texto,
            re.IGNORECASE,
        )
        if m_op_ev:
            datos[_TRIBUTO_ORIGINAL] = m_op_ev.group(1).strip()
            datos[_INTERES_ORIGINAL] = m_op_ev.group(2).strip()
            datos[_TOTAL_ORIGINAL] = m_op_ev.group(3).strip()


def _ficha_seccion_8(texto, datos):
    # 8. Declaración Inicial (Formulario y Documento)
    m_dec = re.search(r"Formulario\s+(\d+)\s+Documento\s+(\d+)", texto, re.IGNORECASE)
    if m_dec:
        # el formulario se escribe con 4 dígitos en los Word: 710 -> 0710, 621 -> 0621
        datos[_FORMULARIO_PDT] = m_dec.group(1).strip().zfill(4)
        datos[_NRO_ORDEN_DECLARACION] = m_dec.group(2).strip()
    else:
        m_dec2 = re.search(r"Formulario\s+621\s+(\d+)", texto, re.IGNORECASE)
        if m_dec2:
            datos[_FORMULARIO_PDT] = "0621"
            datos[_NRO_ORDEN_DECLARACION] = m_dec2.group(1).strip()


def _ficha_seccion_9(texto, datos):
    # 9. Etapas: Exp.Error Material (Inicio = Expediente, Fin = Resolución)
    m_et = re.search(
        r"Exp\.?\s*Error\s+Material\s+(\d{2}/\d{2}/\d{4})\s+(\d{2}/\d{2}/\d{4})\s+\d+\s+(\d+)\s+(\d+)",
        texto,
        re.IGNORECASE,
    )
    if m_et:
        datos[_FECHA_EXPEDIENTE] = m_et.group(1).strip()
        datos[_FECHA_RESOLUCION] = fecha_larga(m_et.group(2).strip())
        datos[_NRO_EXPEDIENTE] = m_et.group(3).strip()
        datos[_NRO_RESOLUCION] = m_et.group(4).strip()
    else:
        m_exp_ev = re.search(
            r"(\d{2}/\d{2}/\d{4})\s+\d{2}:\d{2}\s+Emisi[oó]n\s+de\s+Expediente\s+error\s+material\s+\d+\s+(\d+)",
            texto,
            re.IGNORECASE,
        )
        if m_exp_ev:
            datos[_FECHA_EXPEDIENTE] = m_exp_ev.group(1).strip()
            datos[_NRO_EXPEDIENTE] = m_exp_ev.group(2).strip()


def _ficha_seccion_10(texto, datos):
    # 10. Montos Modificados (Evento Emisión de R. Error Material c Expediente)
    m_rem = re.search(
        r"Emisi[oó]n\s+de\s+R\.?\s*Error\s+Material\s+c\s+Expediente\s+\d+\s+(?:\d+\s+)?([\d,]+\.\d{2})\s+[\d,]+\.\d{2}\s+([\d,]+\.\d{2})\s+([\d,]+\.\d{2})",
        texto,
        re.IGNORECASE,
    )
    if m_rem:
        datos[_TRIBUTO_MODIFICADO] = m_rem.group(1).strip()
        datos[_INTERES_MODIFICADO] = m_rem.group(2).strip()
        datos[_TOTAL_MODIFICADO] = m_rem.group(3).strip()
    else:
        m_mov = re.search(
            r"Ajuste\s+por\s+R\.?\s*Error\s+Material\s+c\s+Expe.*?\(?([\d,]+\.\d{2})\)?\s+[\d,]+\.\d{2}\s+\(?([\d,]+\.\d{2})\)?\s+\(?([\d,]+\.\d{2})\)?",
            texto,
            re.IGNORECASE,
        )
        if m_mov:
            datos[_TRIBUTO_MODIFICADO] = m_mov.group(1).strip()
            datos[_INTERES_MODIFICADO] = m_mov.group(2).strip()
            datos[_TOTAL_MODIFICADO] = m_mov.group(3).strip()


def _ficha_seccion_11(texto, datos):
    # 11. Pagos con Boleta General
    m_pago = re.search(
        r"Pago\s+con\s+Boleta\s+General\s+(\d+)\s+(\d+)\s+([\d,]+\.\d{2})\s+[\d,]+\.\d{2}\s+([\d,]+\.\d{2})\s+([\d,]+\.\d{2})",
        texto,
        re.IGNORECASE,
    )
    if m_pago:
        datos["nro_formulario_pago"] = m_pago.group(1).strip()
        datos[_NRO_BOLETA_PAGO] = m_pago.group(2).strip()
        datos["pago"] = "S/ " + m_pago.group(5).strip()
        m_fec_pago = re.search(
            r"(\d{2}/\d{2}/\d{4})\s+\d{2}:\d{2}\s+Pago\s+con\s+Boleta\s+General", texto, re.IGNORECASE
        )
        if m_fec_pago:
            datos[_FECHA_PAGO] = m_fec_pago.group(1).strip()


def _dec(s):
    s = s.replace(",", "")
    neg = s.startswith("(")
    v = Decimal(s.strip("()"))
    return -v if neg else v


def _fmt(v):
    return f"{v:,.2f}"


def _leer_movimientos_saldo(texto):
    """Filas de 'Movimientos de Saldo': descripción, tributo, interés, total y saldo.

    La celda Tributo puede quedar pisada por la descripción en el PDF, por eso se leen solo los
    últimos 4 montos (Capitalizado, Interés, Total, Saldo) y Tributo = Total - Capitalizado - Interés.
    """
    if not re.search(r"Movimientos\s+de\s+Saldo", texto, re.IGNORECASE):
        return []
    seccion = re.split(r"Movimientos\s+de\s+Saldo", texto, flags=re.IGNORECASE)[-1]
    movs = []
    for linea in seccion.splitlines():
        m_fila = re.match(r"^\d{2}/\d{2}/\d{4}\s+\d{2}:\d{2}\s+(.*)$", linea.strip())
        if not m_fila:
            continue
        cuerpo = re.sub(r"\s+\d+\s+\d{13}\s+\d{2}/\d{2}/\d{4}.*$", "", m_fila.group(1))
        montos = re.findall(r"\(?[\d,]+\.\d{2}\)?", cuerpo)
        if len(montos) < 4:
            continue
        cap, inte, tot, sal = (_dec(x) for x in montos[-4:])
        movs.append({"desc": cuerpo, "trib": tot - cap - inte, "int": inte, "tot": tot, "sal": sal})
    return movs


def _aplicar_ajuste_error_material(datos, movs, ultimo_ajuste):
    # Saldo acumulado tras el ajuste por error material: 0.00 = FUNDADO; cualquier otro valor = FUNDADO EN PARTE
    saldo = movs[ultimo_ajuste]["sal"]
    datos["saldo_error_material"] = _fmt(Decimal("0.00") if saldo == 0 else saldo)
    datos[_FUNDADO] = _FUNDADO_2 if saldo == 0 else "FUNDADO EN PARTE"
    acum = movs[: ultimo_ajuste + 1]
    t_mod = sum(m["trib"] for m in acum)
    i_mod = sum(m["int"] for m in acum)
    tot_mod = sum(m["tot"] for m in acum)
    datos[_MULTA_MOD] = _fmt(t_mod)
    datos[_INTERES_MULTA_MOD] = _fmt(i_mod)
    datos[_TOTAL_MULTA_MOD] = _fmt(tot_mod)
    if tot_mod != acum[-1]["sal"]:
        datos[_AVISOS_EXTRACCION].append(
            "Movimientos de Saldo: el total modificado no coincide con el saldo acumulado; verificar montos."
        )


def _ficha_seccion_10b(texto, datos):
    """Movimientos de Saldo: multa, interés y total (original y modificado).

    - Fila "Deuda por RM"  -> valores originales de la multa.
    - Filas hasta el último "Ajuste por R. Error Material" -> acumulado = valores modificados.
    """
    movs = _leer_movimientos_saldo(texto)
    if not movs:
        datos[_AVISOS_EXTRACCION].append(
            "No se pudo leer Movimientos de Saldo; multa e intereses tomados de otras secciones, verificar."
        )
        return
    orig = next((m for m in movs if re.match(r"Deuda", m["desc"], re.IGNORECASE)), movs[0])
    datos[_MULTA] = _fmt(orig["trib"])
    datos[_INTERES_MULTA] = _fmt(orig["int"])
    datos[_TOTAL_MULTA] = _fmt(orig["tot"])
    idx_aj = [
        i
        for i, m in enumerate(movs)
        if re.match(r"Ajuste\s+por\s+R\.?\s*Error\s+Material", m["desc"], re.IGNORECASE)
    ]
    if idx_aj:
        _aplicar_ajuste_error_material(datos, movs, idx_aj[-1])
    else:
        datos[_AVISOS_EXTRACCION].append(
            "Movimientos de Saldo: no hay 'Ajuste por R. Error Material'; multa modificada tomada de Eventos."
        )


def _completar_multa_desde_valor(datos):
    if datos[_MULTA] == _COMPLETAR:
        datos[_MULTA] = datos[_TRIBUTO_ORIGINAL]
        datos[_INTERES_MULTA] = datos[_INTERES_ORIGINAL]
        datos[_TOTAL_MULTA] = datos[_TOTAL_ORIGINAL]
    if datos[_MULTA_MOD] == _COMPLETAR:
        datos[_MULTA_MOD] = datos[_TRIBUTO_MODIFICADO]
        datos[_INTERES_MULTA_MOD] = datos[_INTERES_MODIFICADO]
        datos[_TOTAL_MULTA_MOD] = datos[_TOTAL_MODIFICADO]


def extraer_datos_pdf(ruta_pdf):
    datos = dict.fromkeys(["archivo", "ruc", _RAZON_SOCIAL, _DOMICILIO, _NRO_RESOLUCION, _CIUDAD, _FECHA_RESOLUCION, _NRO_EXPEDIENTE, _FECHA_EXPEDIENTE, _NRO_ORDEN_PAGO, _PERIODO, _CODIGO_TRIBUTO, _NOMBRE_TRIBUTO, _TRIBUTO_ORIGINAL, _INTERES_ORIGINAL, _TOTAL_ORIGINAL, _FECHA_EMISION_OP, _FECHA_DECLARACION, _NRO_ORDEN_DECLARACION, "fecha_rectificatoria", "nro_orden_rectificatoria", _TRIBUTO_MODIFICADO, _INTERES_MODIFICADO, _TOTAL_MODIFICADO, _FECHA_INTERES, "nro_formulario_pago", _NRO_BOLETA_PAGO, "pago", _FECHA_PAGO, _INTENDENCIA, _CODIGO_TRIBUTO_ASOC, _MULTA, _INTERES_MULTA, _TOTAL_MULTA, _MULTA_MOD, _INTERES_MULTA_MOD, _TOTAL_MULTA_MOD, "saldo_error_material", _FUNDADO, _TIPO_VALOR, _FORMULARIO_PDT], _COMPLETAR)
    datos[_AVISOS_EXTRACCION] = []
    datos["archivo"] = os.path.basename(ruta_pdf)

    try:
        with pdfplumber.open(ruta_pdf) as pdf:
            texto = "\n".join(p.extract_text() or "" for p in pdf.pages[:MAX_PAGINAS_PDF])
    except Exception as e:
        print(f"Error parseando PDF en RIbot: {e}")
        return datos

    if not texto:
        return datos

    _ficha_seccion_1(texto, datos)
    _ficha_seccion_2(texto, datos)
    _ficha_seccion_3(texto, datos)
    _ficha_seccion_4(texto, datos)
    _ficha_seccion_5(texto, datos)
    _ficha_seccion_6(texto, datos)
    _ficha_seccion_7(texto, datos)
    _ficha_seccion_8(texto, datos)
    _ficha_seccion_9(texto, datos)
    _ficha_seccion_10(texto, datos)
    _ficha_seccion_10b(texto, datos)
    _completar_multa_desde_valor(datos)
    _ficha_seccion_11(texto, datos)
    return datos


# ── Generador de Resolución en Word (Carga y Reemplazo desde Plantilla Seleccionada) ──────────────────
PLACEHOLDER = _COMPLETAR
FRASE_FUNDADO = (
    "FUNDADO EN PARTE"  # los modelos la traen escrita; si el saldo del error material es 0.00 pasa a FUNDADO
)
_SOBRA_SI_FUNDADO = " EN PARTE"


def resolver_plantilla(plantillas_dir: str, solicitada: str | None) -> tuple[str, list[str]]:
    """Ruta de la plantilla pedida (sin distinguir mayúsculas) o, si no existe o no se pidió, la primera disponible."""
    avisos: list[str] = []
    existentes = (
        sorted(f for f in os.listdir(plantillas_dir) if f.lower().endswith(_DOCX) and not f.startswith("~$"))
        if os.path.isdir(plantillas_dir)
        else []
    )

    nombre = os.path.basename(solicitada or "")
    if not nombre:
        if existentes:
            return os.path.join(plantillas_dir, existentes[0]), avisos
        raise FileNotFoundError("No hay plantillas Word disponibles")
    if not nombre.lower().endswith(_DOCX):
        nombre += _DOCX

    for f in existentes:
        if f.lower() == nombre.lower():
            return os.path.join(plantillas_dir, f), avisos
    if existentes:
        avisos.append(f"No se encontró la plantilla '{nombre}'; se usó '{existentes[0]}'.")
        return os.path.join(plantillas_dir, existentes[0]), avisos
    raise FileNotFoundError(f"No se encontró la plantilla Word '{nombre}'")


def construir_mapeo(d, nombre_plantilla: str) -> dict:
    """Etiquetas {{CAMPO}} de las plantillas y su valor. Varias etiquetas largas tienen una versión corta."""

    def get(clave, defecto=PLACEHOLDER):
        return str(d.get(clave, defecto))

    # En MODELO INC B y PE, la rectificatoria en los considerandos se refiere al pago de la deuda
    nombre = nombre_plantilla.lower()
    if "inc b" in nombre or "modelo pe" in nombre:
        fecha_recti = d.get(_FECHA_PAGO, "[FECHA PAGO]")
        nro_od_mod = d.get(_NRO_BOLETA_PAGO, "[NRO BOLETA PAGO]")
    else:
        # En MODELO DJR y similares, corresponde a la Declaración Rectificatoria PDT
        f_rec, od_rec = d.get("fecha_rectificatoria"), d.get("nro_orden_rectificatoria")
        fecha_recti = (
            f_rec
            if (f_rec and f_rec != PLACEHOLDER and f_rec != d.get(_FECHA_EXPEDIENTE))
            else "[FECHA RECTIFICATORIA]"
        )
        nro_od_mod = (
            od_rec
            if (od_rec and od_rec != PLACEHOLDER and od_rec != d.get(_NRO_EXPEDIENTE))
            else "[NRO ORDEN RECTIFICATORIA]"
        )

    fec_dec = d.get(_FECHA_DECLARACION)
    if not fec_dec or fec_dec == PLACEHOLDER:
        fec_dec = "[FECHA DECLARACIÓN]"

    return {
        "{{NRO_RES}}": get(_NRO_RESOLUCION),
        "{{CIUDAD}}": get(_CIUDAD, _HUANCAYO),
        "{{FECHA_RES}}": get(_FECHA_RESOLUCION),
        "{{NRO_EXP}}": get(_NRO_EXPEDIENTE),
        "{{FEC_EXP}}": get(_FECHA_EXPEDIENTE),
        "{{RAZON_SOCIAL}}": get(_RAZON_SOCIAL),
        "{{RUC}}": get("ruc"),
        "{{DOMICILIO}}": get(_DOMICILIO),
        "{{DOMICILI}}": get(_DOMICILIO),
        "{{NRO_OP}}": get(_NRO_ORDEN_PAGO),
        "{{COD_TRIB}}": get(_CODIGO_TRIBUTO),
        "{{PERIODO}}": get(_PERIODO),
        "{{FEC_EM_OP}}": get(_FECHA_EMISION_OP),
        "{{TRIB}}": get(_TRIBUTO_ORIGINAL, "0.00"),
        "{{INT}}": get(_INTERES_ORIGINAL, "0.00"),
        "{{TOTAL}}": get(_TOTAL_ORIGINAL, "0.00"),
        "{{COD_TRIB_ASOC}}": get(_CODIGO_TRIBUTO_ASOC),
        "{{CT_ASOC}}": get(_CODIGO_TRIBUTO_ASOC),
        "{{MULTA}}": get(_MULTA, "0.00"),
        "{{INT_MULTA}}": get(_INTERES_MULTA, "0.00"),
        "{{TOTAL_MULTA}}": get(_TOTAL_MULTA, "0.00"),
        "{{T_MULTA}}": get(_TOTAL_MULTA, "0.00"),
        "{{MULTA_MOD}}": get(_MULTA_MOD, "0.00"),
        "{{INT_MULTA_MOD}}": get(_INTERES_MULTA_MOD, "0.00"),
        "{{I_MULTA_M}}": get(_INTERES_MULTA_MOD, "0.00"),
        "{{TOTAL_MULTA_MOD}}": get(_TOTAL_MULTA_MOD, "0.00"),
        "{{T_MULTA_M}}": get(_TOTAL_MULTA_MOD, "0.00"),
        "{{NOM_TRIB_ASOC}}": NOMBRES_TRIBUTO.get(d.get(_CODIGO_TRIBUTO_ASOC), PLACEHOLDER),
        "{{FEC_INT}}": get(_FECHA_INTERES),
        "{{FEC_DEC}}": str(fec_dec),
        "{{NRO_OD}}": get(_NRO_ORDEN_DECLARACION),
        "{{PDT}}": get(_FORMULARIO_PDT),
        "{{FECHA_RECTI}}": str(fecha_recti),
        "{{NRO_OD_MOD}}": str(nro_od_mod),
        "{{TRIB_MOD}}": get(_TRIBUTO_MODIFICADO, "0.00"),
        "{{INT_MOD}}": get(_INTERES_MODIFICADO, "0.00"),
        "{{TOTAL_MOD}}": get(_TOTAL_MODIFICADO, "0.00"),
        "{{PAGO}}": get("pago", "S/ 0.00"),
        "{{FECHA_PE }}": get(_FECHA_EXPEDIENTE),
        "{{RI_PE}}": get(_NRO_RESOLUCION),
        "{{URD}}": get(_NRO_EXPEDIENTE),
        "{{INTENDENCIA}}": get(_INTENDENCIA),
    }


def quitar_en_parte(p) -> bool:
    """Convierte 'FUNDADO EN PARTE' en 'FUNDADO' dentro del párrafo, conservando el formato de cada run.

    La frase puede venir partida en varios runs (p. ej. 'FUNDADO' + ' EN PARTE'): se borran solo los
    caracteres de ' EN PARTE', sin tocar el resto del texto. Devuelve True si hubo algún cambio.
    """
    cambio = False
    while True:
        runs = p.runs
        textos = [r.text or "" for r in runs]
        inicio = "".join(textos).find(FRASE_FUNDADO)
        if inicio == -1:
            return cambio
        inicio += len(_FUNDADO_2)
        fin = inicio + len(_SOBRA_SI_FUNDADO)
        pos = 0
        for run, texto in zip(runs, textos):
            lo, hi = max(pos, inicio), min(pos + len(texto), fin)
            if lo < hi:
                run.text = texto[: lo - pos] + texto[hi - pos :]
            pos += len(texto)
        cambio = True


_MARCADOR = re.compile(r"\{\{[^}]*\}\}")


def _runs_del_marcador(textos, inicio, fin):
    """Índices (primero, último) de los runs que contienen el marcador; -1 si no se hallan."""
    pos, primero, ultimo = 0, -1, -1
    for i, texto in enumerate(textos):
        if pos <= inicio < pos + len(texto):
            primero = i
        if pos < fin <= pos + len(texto):
            return primero, i
        pos += len(texto)
    return primero, ultimo


def _unir_marcadores_partidos(runs, textos) -> None:
    """Une en un solo run los marcadores {{...}} que Word partió en varios runs (actualiza runs y textos)."""
    completo = "".join(textos)
    if "}}" not in completo:
        return
    for m in re.finditer(r"\{\{[^}]+\}\}", completo):
        primero, ultimo = _runs_del_marcador(textos, m.start(), m.end())
        if primero != -1 and ultimo > primero:
            textos[primero] = "".join(textos[primero : ultimo + 1])
            runs[primero].text = textos[primero]
            for j in range(primero + 1, ultimo + 1):
                textos[j] = ""
                runs[j].text = ""


def _sustituir_marcadores(texto: str, mapeo: dict) -> str:
    nuevo = texto
    for marcador in set(_MARCADOR.findall(texto)):
        if marcador in mapeo:
            nuevo = nuevo.replace(marcador, mapeo[marcador])
    return nuevo


def _reemplazo_de_seguridad(p, mapeo: dict) -> None:
    """Sustitución final para llaves que queden fragmentadas entre runs."""
    completo = "".join(r.text or "" for r in p.runs)
    if "{{" not in completo:
        return
    reemplazado = _sustituir_marcadores(completo, mapeo)
    if reemplazado == completo:
        return
    for r in p.runs:
        r.text = ""
    if p.runs:
        p.runs[0].text = reemplazado
    else:
        p.add_run(reemplazado)


def reemplazar_en_parrafo(p, mapeo: dict) -> None:
    # Cada acceso a run.text recorre el XML: se lee una sola vez y solo se trabaja si hay marcadores
    runs = p.runs
    textos = [r.text or "" for r in runs]
    if "{{" not in "".join(textos):
        return
    _unir_marcadores_partidos(runs, textos)

    for run, texto in zip(runs, textos):
        if "{{" not in texto:
            continue
        nuevo = _sustituir_marcadores(texto, mapeo)
        if nuevo != texto:
            run.text = nuevo

    _reemplazo_de_seguridad(p, mapeo)


def _parrafos_de_tablas(tablas):
    for tabla in tablas:
        for fila in tabla.rows:
            for celda in fila.cells:
                yield from celda.paragraphs


def todos_los_parrafos(doc):
    yield from doc.paragraphs
    yield from _parrafos_de_tablas(doc.tables)
    for seccion in doc.sections:
        for parte in (seccion.header, seccion.footer):
            if parte:
                yield from parte.paragraphs
                yield from _parrafos_de_tablas(parte.tables)


def _avisos_de_plantilla(doc, d, tipo, habia_frase, verbo_aplicado) -> list:
    """Avisos por campos sin dato, llaves desconocidas o decisiones que el colaborador debe verificar."""
    avisos = []
    cuerpo = list(doc.paragraphs) + list(_parrafos_de_tablas(doc.tables))
    texto_final = "\n".join(p.text for p in cuerpo)
    pendientes = sorted(set(re.findall(r"\[[A-ZÁÉÍÓÚ ]{4,}\]", texto_final)))
    if pendientes:
        avisos.append("Datos no encontrados en el PDF: " + ", ".join(pendientes) + ".")
    sin_reemplazar = sorted(set(re.findall(r"\{\{[^}]*\}\}", texto_final)) - {TEXTO_LIBRE_MARCADOR})
    if sin_reemplazar:
        avisos.append("Campos de la plantilla sin reemplazar: " + ", ".join(sin_reemplazar) + ".")

    # Sin la fila de ajuste en la Ficha no se puede decidir: el modelo queda con su texto y se avisa
    if d.get(_FUNDADO) in (None, "", PLACEHOLDER) and FRASE_FUNDADO in texto_final:
        avisos.append(
            "No se encontró el ajuste por error material en Movimientos de Saldo: verifica si corresponde "
            "FUNDADO o FUNDADO EN PARTE (el modelo quedó con FUNDADO EN PARTE)."
        )

    if habia_frase and not verbo_aplicado and FRASE_FUNDADO not in texto_final:
        avisos.append(
            "No se encontró el ajuste por error material en Movimientos de Saldo: verifica si corresponde "
            "'modificar' o 'revocar' en el modelo."
        )
    if habia_frase and tipo is None:
        avisos.append("No se encontró el tipo de documento (Orden de Pago, Resolución de Multa...) en la Ficha.")
    return avisos


def generar_word(d, ruta_salida, nombre_plantilla=None, texto_libre=""):
    plantillas_dir = obtener_directorio_plantillas()
    ruta_plantilla, avisos = resolver_plantilla(plantillas_dir, nombre_plantilla)
    avisos = list(d.get(_AVISOS_EXTRACCION, [])) + avisos

    doc = Document(ruta_plantilla)
    mapeo = construir_mapeo(d, os.path.basename(ruta_plantilla))
    parrafos = list(todos_los_parrafos(doc))
    for p in parrafos:
        reemplazar_en_parrafo(p, mapeo)

    # Saldo acumulado del ajuste por error material = 0.00 -> el modelo dice FUNDADO (no "FUNDADO EN PARTE")
    if d.get(_FUNDADO) == _FUNDADO_2:
        for p in parrafos:
            if FRASE_FUNDADO in p.text:
                quitar_en_parte(p)

    # "modificar/revocar la orden de pago": el verbo sale del FUNDADO y el tipo, de la línea Valor de la Ficha
    tipo = d.get(_TIPO_VALOR)
    tipo = None if tipo in (None, "", PLACEHOLDER) else tipo
    habia_frase, verbo_aplicado = adapt_resolutivo(parrafos, d.get(_FUNDADO), tipo)

    # Se calculan antes de insertar el texto libre para no confundir lo que escribió el colaborador.
    avisos += _avisos_de_plantilla(doc, d, tipo, habia_frase, verbo_aplicado)

    if not aplicar_texto_libre(parrafos, texto_libre) and normalizar_lineas(texto_libre):
        avisos.append(
            f"El modelo '{os.path.splitext(os.path.basename(ruta_plantilla))[0]}' no admite texto adicional; no se incluyó."
        )

    doc.save(ruta_salida)
    return avisos


def completar_desde_padron(d, etiqueta: str, lote: dict) -> None:
    """Completa domicilio fiscal (y razón social si falta) consultando el RUC en el web service interno."""
    cliente = get_padron()
    if lote.get(_PADRON_CAIDO) or not cliente.enabled:
        return
    if d.get(_DOMICILIO) not in (None, "", PLACEHOLDER):
        return
    if d.get("ruc") in (None, "", PLACEHOLDER):
        lote[_AVISOS].append(f"{etiqueta}: no se encontró el RUC en el PDF; no se pudo consultar el domicilio fiscal.")
        return
    try:
        registro = cliente.consultar(d["ruc"])
    except PadronUnavailableError as exc:
        # No se reintenta con el resto del lote para no esperar el timeout en cada PDF
        lote[_PADRON_CAIDO] = True
        lote[_AVISOS].append(
            f"No se pudo consultar el domicilio en el padrón SUNAT: {exc}. Se omite para el resto del lote."
        )
        return
    except PadronError as exc:
        lote[_AVISOS].append(f"{etiqueta}: no se obtuvo el domicilio del RUC {d['ruc']} ({exc}).")
        return
    if registro.domicilio:
        d[_DOMICILIO] = formato_domicilio(registro.domicilio)
    if registro.razon_social and d.get(_RAZON_SOCIAL) in (None, "", PLACEHOLDER):
        d[_RAZON_SOCIAL] = registro.razon_social


def texto_para(archivo: str, general: str, por_archivo: dict) -> str:
    """Redacción de la evaluación de un PDF: la suya, si la tiene, o la general."""
    propio = por_archivo.get(archivo, "")
    return propio if str(propio).strip() else general


def nombre_unico(nombre: str, usados: set) -> str:
    """Evita que dos PDFs del mismo lote (p. ej. ambos sin datos) se pisen entre sí."""
    base, ext = os.path.splitext(nombre)
    candidato, n = nombre, 2
    while candidato.lower() in usados:
        candidato = f"{base}_{n}{ext}"
        n += 1
    usados.add(candidato.lower())
    return candidato


def _json_dict(raw: str) -> dict:
    """Diccionario enviado como JSON en un campo de formulario; {} si viene vacío o inválido."""
    try:
        data = json.loads(raw) if raw and raw.strip() else {}
    except Exception:
        return {}
    return {str(k): str(v) for k, v in data.items()} if isinstance(data, dict) else {}


# ── Endpoints de Gestión de Plantillas / Modelos Word ────────────────────────────


@router.get("/api/models")
async def list_models():
    """Devuelve la lista de plantillas .docx disponibles."""
    plantillas_dir = obtener_directorio_plantillas()
    modelos = []
    if os.path.exists(plantillas_dir):
        for file in os.listdir(plantillas_dir):
            if file.lower().endswith(_DOCX) and not file.startswith("~$"):
                full_path = os.path.join(plantillas_dir, file)
                size_bytes = os.path.getsize(full_path)
                mtime = os.path.getmtime(full_path)
                nombre_display = os.path.splitext(file)[0]
                modelos.append({_FILENAME: file, "name": nombre_display, "size": size_bytes, "mtime": mtime})
    modelos.sort(key=lambda x: x["name"].lower())
    return {_STATUS: _SUCCESS, "models": modelos}


@router.post("/api/open-templates-folder")
async def open_templates_folder():
    """Abre la carpeta física de plantillas Word en el Explorador de Windows."""
    try:
        plantillas_dir = obtener_directorio_plantillas()
        if os.name == "nt" and os.path.exists(plantillas_dir):
            os.startfile(plantillas_dir)
        return {_STATUS: _SUCCESS, "path": plantillas_dir}
    except Exception as e:
        return JSONResponse(status_code=500, content={_STATUS: _ERROR, _MESSAGE: str(e)})


@router.post("/api/upload-model")
async def upload_model(files: List[UploadFile] = File(...)):
    """Sube una o varias plantillas .docx al directorio de plantillas de RIbot."""
    if not files:
        return JSONResponse(status_code=400, content={_STATUS: _ERROR, _MESSAGE: "No se enviaron archivos."})

    plantillas_dir = obtener_directorio_plantillas()
    os.makedirs(plantillas_dir, exist_ok=True)
    subidos = []
    for file in files:
        if not file.filename.lower().endswith(_DOCX):
            continue
        nombre_limpio = sanitizar_nombre_archivo(os.path.splitext(file.filename)[0]) + _DOCX
        dest_path = os.path.join(plantillas_dir, nombre_limpio)
        with open(dest_path, "wb") as f:
            shutil.copyfileobj(file.file, f)
        subidos.append(nombre_limpio)

    if not subidos:
        return JSONResponse(
            status_code=400,
            content={_STATUS: _ERROR, _MESSAGE: "Formato inválido. Solo se admiten plantillas .docx"},
        )

    return JSONResponse(
        content={
            _STATUS: _SUCCESS,
            _MESSAGE: f"Se agregaron {len(subidos)} plantilla(s) correctamente.",
            "uploaded": subidos,
        }
    )


@router.delete("/api/delete-model/{filename}")
async def delete_model(filename: str):
    """Elimina una plantilla personalizada .docx."""
    safe_name = os.path.basename(filename)
    plantillas_dir = obtener_directorio_plantillas()
    dest_path = os.path.join(plantillas_dir, safe_name)
    if os.path.exists(dest_path):
        os.remove(dest_path)
        return JSONResponse(
            content={_STATUS: _SUCCESS, _MESSAGE: f"Plantilla '{safe_name}' eliminada correctamente."}
        )
    return JSONResponse(status_code=404, content={_STATUS: _ERROR, _MESSAGE: "Plantilla no encontrada."})


# ── Endpoints de la API del Módulo RIbot ───────────────────────────────────────────


@router.get("/")
async def ribot_home(request: Request):
    limpiar_archivos_temporales()
    return templates.TemplateResponse("ribot/index.html", {"request": request, "modulo": "RIbot"})


def _error_json(message, status_code=400):
    return JSONResponse(status_code=status_code, content={_STATUS: _ERROR, _MESSAGE: message})


def _texto_excedido(texto_libre, textos_por_pdf):
    """Respuesta 422 si alguna redacción supera TEXTO_MAX; None si todo está bien."""
    if len(texto_libre) > TEXTO_MAX or any(len(str(t)) > TEXTO_MAX for t in textos_por_pdf.values()):
        return _error_json(f"La redacción de la evaluación supera {TEXTO_MAX} caracteres.", 422)
    return None


def _crear_destino(destination_path):
    """Crea la carpeta de destino. Devuelve (ruta o None, respuesta de error o None)."""
    if not destination_path.strip():
        return None, None
    try:
        dest_dir = os.path.abspath(destination_path.strip())
        os.makedirs(dest_dir, exist_ok=True)
        return dest_dir, None
    except Exception as e:
        return None, _error_json(f"Ruta de destino inválida o sin permisos: {e}")


def _carpeta_origen(carpeta):
    """Resuelve la carpeta de PDFs. Devuelve (ruta o None, respuesta de error o None)."""
    is_absolute = os.path.isabs(carpeta) or (len(carpeta) >= 2 and carpeta[1] == ":")
    try:
        if is_absolute:
            target_path = os.path.abspath(carpeta)
        else:
            target_path = os.path.abspath(os.path.join(RIBOT_BASE_DIR, carpeta.strip(r"\/")))
            if not target_path.startswith(os.path.abspath(RIBOT_BASE_DIR)):
                return None, _error_json("Acceso denegado a esta ruta")
    except Exception as e:
        return None, _error_json(f"Ruta inválida: {e}")

    if not os.path.isdir(target_path):
        return None, _error_json("La carpeta de origen no existe en el servidor.")
    return target_path, None


async def _extraer_y_nombrar(pdf_path, etiqueta, lote, usados):
    """Extrae los datos del PDF, los completa con el padrón y arma el nombre único del Word."""
    d = await run_in_threadpool(extraer_datos_pdf, pdf_path)
    await run_in_threadpool(completar_desde_padron, d, etiqueta, lote)
    ruc_safe = sanitizar_nombre_archivo(d.get("ruc", "SIN_RUC"))
    nro_res_safe = sanitizar_nombre_archivo(d.get(_NRO_RESOLUCION, "SIN_NRO"))
    razon_safe = sanitizar_nombre_archivo(d.get(_RAZON_SOCIAL, "SIN_NOMBRE"))
    return d, nombre_unico(f"{ruc_safe}_{nro_res_safe}_{razon_safe}_mod.docx", usados)


async def _redactar_word(d, ruta_word, modelo, texto, etiqueta, avisos, despues=None):
    """Genera el Word. Devuelve True si quedó listo; si falla, deja el aviso y devuelve False."""
    try:
        for av in await run_in_threadpool(generar_word, d, ruta_word, modelo, texto):
            avisos.append(f"{etiqueta}: {av}")
        if despues:
            despues()
        return True
    except Exception as e:
        print(f"Error redactando Word {os.path.basename(ruta_word)}: {e}")
        avisos.append(f"{etiqueta}: no se pudo generar el Word ({e}).")
        return False


@router.post("/api/process-upload")
async def process_upload(
    files: List[UploadFile] = File(...),
    destination_path: str = Form(""),
    model_name: str = Form(""),
    file_models: str = Form("{}"),
    texto_libre: str = Form(""),
    file_textos: str = Form("{}"),
):
    """Sube PDFs de Fichas de Valor, genera Word y guarda en destino físico y/o descarga."""
    limpiar_archivos_temporales()

    if not files:
        return JSONResponse(status_code=400, content={_STATUS: _ERROR, _MESSAGE: "No se subieron archivos."})

    dest_dir, error = _crear_destino(destination_path)
    if error:
        return error

    job_id = str(uuid.uuid4())
    job_path = os.path.join(TEMP_DIR, job_id)
    os.makedirs(job_path, exist_ok=True)

    input_path = os.path.join(job_path, "input")
    os.makedirs(input_path, exist_ok=True)

    generated_files = []
    lote = {_AVISOS: [], _PADRON_CAIDO: False}
    avisos = lote[_AVISOS]
    usados = set()
    mapping_modelos = _json_dict(file_models)
    textos_por_pdf = _json_dict(file_textos)
    error = _texto_excedido(texto_libre, textos_por_pdf)
    if error:
        return error

    for file in files:
        if not file.filename.lower().endswith(".pdf"):
            avisos.append(f"{file.filename}: se omitió (no es PDF).")
            continue

        safe_pdf_name = sanitizar_nombre_archivo(os.path.basename(file.filename)) + ".pdf"
        pdf_path = os.path.join(input_path, safe_pdf_name)

        with open(pdf_path, "wb") as f:
            shutil.copyfileobj(file.file, f)

        # Extracción y Redacción
        d, nombre_word = await _extraer_y_nombrar(pdf_path, file.filename, lote, usados)

        # Obtener el modelo asignado para este PDF (o el global; vacío = la primera plantilla)
        modelo_especifico = mapping_modelos.get(file.filename) or model_name
        ruta_word = os.path.join(job_path, nombre_word)
        texto = texto_para(file.filename, texto_libre, textos_por_pdf)

        def copiar_a_destino(ruta_word=ruta_word, nombre_word=nombre_word):
            generated_files.append({_FILENAME: nombre_word, "path": ruta_word})
            if dest_dir:
                shutil.copy2(ruta_word, os.path.join(dest_dir, nombre_word))

        await _redactar_word(d, ruta_word, modelo_especifico, texto, file.filename, avisos, copiar_a_destino)

    if not generated_files:
        return JSONResponse(
            status_code=400,
            content={_STATUS: _ERROR, _MESSAGE: "No se pudieron procesar las Fichas del Valor PDF subidas."},
        )

    if len(generated_files) == 1:
        filename = generated_files[0][_FILENAME]
        download_url = f"/ribot/api/download/{job_id}/{filename}"
        return JSONResponse(
            content={
                _STATUS: _SUCCESS,
                "type": "single",
                _FILENAME: filename,
                _DOWNLOAD_URL: download_url,
                _COUNT: 1,
                _PHYSICAL_PATH: dest_dir,
                _SAVED_LOCALLY: bool(dest_dir),
                _AVISOS: avisos,
            }
        )
    zip_filename = f"Resoluciones_SUNAT_{job_id[:8]}.zip"
    zip_path = os.path.join(job_path, zip_filename)

    with zipfile.ZipFile(zip_path, "w") as zipf:
        for item in generated_files:
            zipf.write(item["path"], item[_FILENAME])

    download_url = f"/ribot/api/download/{job_id}/{zip_filename}"
    return JSONResponse(
        content={
            _STATUS: _SUCCESS,
            "type": "massive",
            _FILENAME: zip_filename,
            _DOWNLOAD_URL: download_url,
            _COUNT: len(generated_files),
            _PHYSICAL_PATH: dest_dir,
            _SAVED_LOCALLY: bool(dest_dir),
            _AVISOS: avisos,
        }
    )


@router.post("/api/process-folder")
async def process_folder(data: dict):
    """Procesa una carpeta del servidor que contiene Fichas de Valor PDF y guarda en destino físico."""
    limpiar_archivos_temporales()

    carpeta = data.get("carpeta", "").strip()
    destination_path = data.get("destination_path", "").strip()
    model_name = data.get("model_name", "")
    mapping_modelos = data.get("file_models", {})
    texto_libre = str(data.get("texto_libre", ""))
    textos_por_pdf = data.get("file_textos", {})
    error = _texto_excedido(texto_libre, textos_por_pdf)
    if error:
        return error

    target_path, error = _carpeta_origen(carpeta)
    if error:
        return error

    try:
        pdfs = [f for f in os.listdir(target_path) if f.lower().endswith(".pdf")]
    except Exception as e:
        return JSONResponse(
            status_code=400, content={_STATUS: _ERROR, _MESSAGE: f"Sin permisos o error al leer la carpeta: {e}"}
        )

    if not pdfs:
        return JSONResponse(
            status_code=400,
            content={
                _STATUS: _ERROR,
                _MESSAGE: "No se encontraron PDFs de Fichas de Valor en la carpeta seleccionada.",
            },
        )

    dest_dir, error = _crear_destino(destination_path)
    if error:
        return error

    job_id = str(uuid.uuid4())
    job_path = os.path.join(TEMP_DIR, job_id)
    os.makedirs(job_path, exist_ok=True)

    generated_files = []
    lote = {_AVISOS: [], _PADRON_CAIDO: False}
    avisos = lote[_AVISOS]
    usados = set()

    for pdf_name in pdfs:
        pdf_path = os.path.join(target_path, pdf_name)
        d, nombre_word = await _extraer_y_nombrar(pdf_path, pdf_name, lote, usados)

        modelo_pdf = mapping_modelos.get(pdf_name) or model_name
        ruta_word = os.path.join(dest_dir or target_path, nombre_word)
        texto = texto_para(pdf_name, texto_libre, textos_por_pdf)

        def registrar(ruta_word=ruta_word, nombre_word=nombre_word):
            generated_files.append({_FILENAME: nombre_word, "path": ruta_word})

        await _redactar_word(d, ruta_word, modelo_pdf, texto, pdf_name, avisos, registrar)

    if not generated_files:
        return JSONResponse(
            status_code=500,
            content={_STATUS: _ERROR, _MESSAGE: "No se pudieron redactar las resoluciones de la carpeta."},
        )

    zip_filename = f"Resoluciones_Local_{job_id[:8]}.zip"
    zip_path = os.path.join(job_path, zip_filename)

    with zipfile.ZipFile(zip_path, "w") as zipf:
        for item in generated_files:
            zipf.write(item["path"], item[_FILENAME])

    download_url = f"/ribot/api/download/{job_id}/{zip_filename}"
    return JSONResponse(
        content={
            _STATUS: _SUCCESS,
            "type": "massive",
            _FILENAME: zip_filename,
            _DOWNLOAD_URL: download_url,
            _COUNT: len(generated_files),
            _PHYSICAL_PATH: dest_dir or target_path,
            _SAVED_LOCALLY: True,
            _AVISOS: avisos,
        }
    )


@router.get("/api/download/{job_id}/{filename}")
async def download_results(job_id: str, filename: str):
    safe_filename = os.path.basename(filename)
    file_path = os.path.join(TEMP_DIR, job_id, safe_filename)
    if os.path.exists(file_path):
        return FileResponse(file_path, filename=safe_filename)
    return JSONResponse(status_code=404, content={_MESSAGE: "Archivo no encontrado"})


def _ruta_de_navegacion(path_str):
    """Ruta absoluta a listar. Devuelve (ruta, es_absoluta, error o None)."""
    if os.path.isabs(path_str) or (len(path_str) >= 2 and path_str[1] == ":"):
        return os.path.abspath(path_str), True, None
    target_path = os.path.abspath(os.path.join(RIBOT_BASE_DIR, path_str.strip(r"\/")))
    if not target_path.startswith(os.path.abspath(RIBOT_BASE_DIR)):
        return None, False, "Acceso denegado"
    return target_path, False, None


def _subcarpetas(target_path, is_absolute):
    folders = []
    for item in os.listdir(target_path):
        full_path = os.path.join(target_path, item)
        if os.path.isdir(full_path):
            ruta = full_path if is_absolute else os.path.relpath(full_path, RIBOT_BASE_DIR)
            folders.append({"name": item, "path": ruta})
    folders.sort(key=lambda x: x["name"].lower())
    return folders


def _ruta_a_mostrar(target_path, is_absolute):
    if is_absolute:
        return target_path
    if target_path != os.path.abspath(RIBOT_BASE_DIR):
        return os.path.relpath(target_path, RIBOT_BASE_DIR)
    return ""


@router.get("/api/list-folders")
async def list_folders(path: str = ""):
    """Lista subcarpetas locales. Soporta rutas absolutas locales o relativas a ribot_docs."""
    try:
        target_path, is_absolute, error = _ruta_de_navegacion(path.strip())
        if error:
            return {"folders": [], "error": error}
        if not os.path.isdir(target_path):
            return {"folders": [], "error": "La ruta especificada no existe o no es un directorio válido."}
        try:
            folders = _subcarpetas(target_path, is_absolute)
        except PermissionError:
            return {"folders": [], "error": "Sin permisos de acceso en esta carpeta."}
        return {
            "folders": folders,
            "current": _ruta_a_mostrar(target_path, is_absolute),
            "is_absolute": is_absolute,
        }
    except Exception as e:
        return {"folders": [], "error": str(e)}

