"""
CSV2Excel Processor - Motor de procesamiento, limpieza y diseño ejecutable de CSV a Excel.
Procesa lotes masivos de 30+ archivos preservando formatos nativos de Excel y creando Tablas Nativa (Ctrl + T).
"""

import io
import os
import re
import csv
import zipfile
from typing import List, Dict, Any, Tuple, Optional
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

# Temas de color predefinidos para las tablas de Excel
THEMES = {
    "excel_green": {
        "name": "Verde Excel Nativo",
        "header_fill": "107C41",
        "header_font": "FFFFFF",
        "zebra_fill": "F2F9F4",
        "border_color": "D0E1D4",
        "accent_color": "0E6233",
        "table_style": "TableStyleMedium9",
    },
    "corporate_blue": {
        "name": "Azul Corporativo SICODE",
        "header_fill": "0163AC",
        "header_font": "FFFFFF",
        "zebra_fill": "F0F6FA",
        "border_color": "D0E2F1",
        "accent_color": "014E88",
        "table_style": "TableStyleMedium2",
    },
    "dark_slate": {
        "name": "Gris Ejecutivo Dark",
        "header_fill": "1E293B",
        "header_font": "FFFFFF",
        "zebra_fill": "F8FAFC",
        "border_color": "E2E8F0",
        "accent_color": "0F172A",
        "table_style": "TableStyleMedium1",
    },
    "sunat_red": {
        "name": "Rojo Institucional SUNAT",
        "header_fill": "B91C1C",
        "header_font": "FFFFFF",
        "zebra_fill": "FEF2F2",
        "border_color": "FCA5A5",
        "accent_color": "991B1B",
        "table_style": "TableStyleMedium3",
    },
}


class CSV2ExcelProcessor:
    def __init__(self):
        pass

    def detect_encoding_and_delimiter(self, raw_bytes: bytes) -> Tuple[str, str]:
        """
        Detecta la codificación (utf-8-sig, utf-8, cp1252, latin-1)
        y el separador (, ; \t |) de un buffer de bytes CSV.
        """
        content_text = ""
        used_encoding = "utf-8"

        for enc in ["utf-8-sig", "utf-8", "cp1252", "latin-1"]:
            try:
                content_text = raw_bytes[:20000].decode(enc)
                used_encoding = enc
                break
            except (UnicodeDecodeError, AttributeError):
                continue

        if not content_text:
            content_text = raw_bytes[:20000].decode("latin-1", errors="ignore")
            used_encoding = "latin-1"

        lines = [line.strip() for line in content_text.splitlines() if line.strip()]
        sample_lines = lines[:15]

        candidates = [";", ",", "\t", "|"]
        scores = {c: 0 for c in candidates}

        for delim in candidates:
            counts = [line.count(delim) for line in sample_lines]
            if counts and counts[0] > 0 and len(set(counts)) == 1:
                scores[delim] += counts[0] * 10
            elif counts:
                scores[delim] += sum(counts)

        best_delimiter = max(scores, key=scores.get)
        if scores[best_delimiter] == 0:
            best_delimiter = ","

        return used_encoding, best_delimiter

    def parse_csv_file(self, raw_bytes: bytes, filename: str, custom_delimiter: Optional[str] = None) -> Dict[str, Any]:
        """
        Parsea un archivo CSV en memoria y retorna DataFrame + metadatos + previsualización.
        """
        encoding, detected_delimiter = self.detect_encoding_and_delimiter(raw_bytes)
        delimiter = custom_delimiter if custom_delimiter else detected_delimiter

        try:
            df = pd.read_csv(
                io.BytesIO(raw_bytes), sep=delimiter, encoding=encoding, dtype=str, on_bad_lines="skip", engine="python"
            )
        except Exception:
            try:
                text = raw_bytes.decode(encoding, errors="ignore")
                reader = csv.reader(io.StringIO(text), delimiter=delimiter)
                rows = list(reader)
                df = pd.DataFrame(rows[1:], columns=rows[0]) if rows else pd.DataFrame()
            except Exception as e2:
                raise ValueError(f"No se pudo analizar el archivo '{filename}': {str(e2)}") from e2

        df.columns = [
            str(c).strip() if pd.notna(c) and str(c).strip() != "" else f"Columna_{i + 1}"
            for i, c in enumerate(df.columns)
        ]
        df = df.dropna(how="all").reset_index(drop=True)
        df = df.fillna("")

        column_info = []
        for col in df.columns:
            sample_vals = [str(v).strip() for v in df[col].tolist() if pd.notna(v) and str(v).strip() != ""][:100]
            col_type = self._infer_column_type(col, sample_vals)
            column_info.append({"name": col, "type": col_type, "sample": sample_vals[:3]})

        preview_rows = df.head(10).to_dict(orient="records")

        return {
            "filename": filename,
            "encoding": encoding,
            "delimiter": delimiter,
            "rows_count": len(df),
            "cols_count": len(df.columns),
            "columns": column_info,
            "preview": preview_rows,
            "df": df,
        }

    def _infer_column_type(self, col_name: str, samples: List[str]) -> str:
        """
        Deduce si la columna es Identificador (RUC/DNI/Doc), Moneda, Fecha, Número o Texto.
        """
        col_lower = col_name.lower()

        id_terms = [
            "ruc",
            "dni",
            "codigo",
            "código",
            "num_",
            "documento",
            "cuenta",
            "telefono",
            "teléfono",
            "ubigeo",
            "placa",
            "guia",
            "guía",
            "ticket",
            "item",
            "clave",
            "id",
            "partida",
            "recibo",
            "factura",
            "boleta",
            "siaf",
            "secuencia",
        ]
        if any(term in col_lower for term in id_terms):
            return "text_id"

        non_empty = [s for s in samples if s != ""]
        if not non_empty:
            return "text"

        has_leading_zero = any(
            (s.startswith("0") and len(s) > 1) or ("-" in s and s.replace("-", "").isdigit()) for s in non_empty
        )
        if has_leading_zero:
            return "text_id"

        is_currency = any(re.match(r"^S/|S/\.|[$€]\s?[\d.,]+$", s) for s in non_empty)
        if is_currency or any(
            term in col_lower
            for term in ["monto", "deuda", "total", "saldo", "precio", "importe", "igv", "valor", "pim", "devengado"]
        ):
            return "currency"

        is_date = any(re.match(r"^\d{1,4}[/-]\d{1,2}[/-]\d{1,4}$", s) for s in non_empty)
        if is_date or "fecha" in col_lower or "periodo" in col_lower or "fec_" in col_lower:
            return "date"

        is_numeric = all(re.match(r"^-?\d+([.,]\d+)?$", s.replace(" ", "")) for s in non_empty)
        if is_numeric:
            return "number"

        return "text"

    def generate_excel_workbook(
        self,
        files_data: List[Dict[str, Any]],
        output_mode: str = "individual_zip",
        theme_key: str = "excel_green",
        preserve_leading_zeros: bool = True,
        add_totals_row: bool = True,
        auto_fit_columns: bool = True,
        freeze_header: bool = True,
        add_index_sheet: bool = True,
    ) -> Tuple[bytes, str, str]:
        """
        Genera el archivo descargable (.xlsx o .zip) aplicando el formato de Tabla Nativa (Ctrl + T).
        """
        theme = THEMES.get(theme_key, THEMES["excel_green"])

        if output_mode == "individual_zip":
            zip_buffer = io.BytesIO()
            existing_filenames = set()

            with zipfile.ZipFile(zip_buffer, "w", zipfile.ZIP_DEFLATED) as zf:
                for idx, f_item in enumerate(files_data):
                    wb = openpyxl.Workbook()
                    wb.remove(wb.active)

                    sheet_name = self._sanitize_sheet_name(f_item["filename"], idx + 1, set())
                    ws = wb.create_sheet(title=sheet_name)

                    self._format_worksheet(
                        ws,
                        f_item["df"],
                        theme,
                        preserve_leading_zeros=preserve_leading_zeros,
                        add_totals_row=add_totals_row,
                        auto_fit_columns=auto_fit_columns,
                        freeze_header=freeze_header,
                        sheet_index=idx + 1,
                    )

                    out_stream = io.BytesIO()
                    wb.save(out_stream)

                    base_name = os.path.splitext(f_item["filename"])[0]
                    clean_fname = re.sub(r"[\:\/\?\*\[\]\\]", "_", base_name).strip() + "_excel.xlsx"

                    dup_cnt = 1
                    final_zip_entry_name = clean_fname
                    while final_zip_entry_name.lower() in existing_filenames:
                        final_zip_entry_name = f"{base_name}_{dup_cnt}_excel.xlsx"
                        dup_cnt += 1
                    existing_filenames.add(final_zip_entry_name.lower())

                    zf.writestr(final_zip_entry_name, out_stream.getvalue())

            return zip_buffer.getvalue(), f"Lote_{len(files_data)}_Excels_Separados.zip", "application/zip"

        if output_mode == "merged_single":
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Consolidado Master"

            combined_dfs = []
            for f_item in files_data:
                df_temp = f_item["df"].copy()
                df_temp.insert(0, "Archivo Origen", f_item["filename"])
                combined_dfs.append(df_temp)

            merged_df = pd.concat(combined_dfs, ignore_index=True) if combined_dfs else pd.DataFrame()

            self._format_worksheet(
                ws,
                merged_df,
                theme,
                preserve_leading_zeros=preserve_leading_zeros,
                add_totals_row=add_totals_row,
                auto_fit_columns=auto_fit_columns,
                freeze_header=freeze_header,
                sheet_index=1,
            )

            out_stream = io.BytesIO()
            wb.save(out_stream)
            return (
                out_stream.getvalue(),
                f"Consolidado_Master_{len(files_data)}_Archivos.xlsx",
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )

        wb = openpyxl.Workbook()
        wb.remove(wb.active)

        sheets_summary = []
        used_sheet_names = set()

        for idx, f_item in enumerate(files_data):
            sheet_name = self._sanitize_sheet_name(f_item["filename"], idx + 1, used_sheet_names)
            ws = wb.create_sheet(title=sheet_name)

            self._format_worksheet(
                ws,
                f_item["df"],
                theme,
                preserve_leading_zeros=preserve_leading_zeros,
                add_totals_row=add_totals_row,
                auto_fit_columns=auto_fit_columns,
                freeze_header=freeze_header,
                sheet_index=idx + 1,
            )

            sheets_summary.append(
                {
                    "sheet_name": sheet_name,
                    "filename": f_item["filename"],
                    "rows": len(f_item["df"]),
                    "cols": len(f_item["df"].columns),
                }
            )

        if add_index_sheet and len(sheets_summary) > 0:
            self._create_index_sheet(wb, sheets_summary, theme)

        out_stream = io.BytesIO()
        wb.save(out_stream)
        return (
            out_stream.getvalue(),
            f"Lote_{len(files_data)}_CSVs_MultiPestana.xlsx",
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

    def _sanitize_sheet_name(self, filename: str, fallback_idx: int, existing_names: set) -> str:
        base = os.path.splitext(filename)[0]
        cleaned = re.sub(r"[\:\/\?\*\[\]\\]", "_", base).strip()
        if not cleaned:
            cleaned = f"Hoja_{fallback_idx}"

        cleaned = cleaned[:27]
        final_name = cleaned
        counter = 1
        while final_name.lower() in [n.lower() for n in existing_names]:
            final_name = f"{cleaned}_{counter}"
            counter += 1

        existing_names.add(final_name)
        return final_name

    def _create_index_sheet(self, wb: openpyxl.Workbook, summary_list: List[Dict[str, Any]], theme: Dict[str, str]):
        ws = wb.create_sheet(title="📌 Índice", index=0)
        ws.views.sheetView[0].showGridLines = True

        ws.merge_cells("B2:F2")
        title_cell = ws["B2"]
        title_cell.value = "📊 RESUMEN GENERAL DEL LOTE PROCESADO"
        title_cell.font = Font(name="Segoe UI", size=15, bold=True, color=theme["accent_color"])
        title_cell.alignment = Alignment(vertical="center")

        ws.merge_cells("B3:F3")
        sub_cell = ws["B3"]
        sub_cell.value = (
            f"Total de archivos procesados: {len(summary_list)} • Haz clic para ir a la pestaña correspondiente"
        )
        sub_cell.font = Font(name="Segoe UI", size=10, italic=True, color="64748B")

        headers = ["N°", "Nombre de Pestaña", "Archivo Origen", "Total Filas", "Total Columnas", "Acceso Directo"]
        header_fill = PatternFill(fgColor=theme["header_fill"], fill_type="solid")
        header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")

        start_row = 5
        for col_idx, h_text in enumerate(headers, start=2):
            cell = ws.cell(row=start_row, column=col_idx, value=h_text)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center" if col_idx != 3 else "left", vertical="center")

        border_thin = Side(style="thin", color="CBD5E1")
        border_all = Border(left=border_thin, right=border_thin, top=border_thin, bottom=border_thin)

        for i, item in enumerate(summary_list, start=1):
            row_idx = start_row + i
            is_zebra = i % 2 == 0
            row_fill = PatternFill(fgColor=theme["zebra_fill"], fill_type="solid") if is_zebra else None

            c_num = ws.cell(row=row_idx, column=2, value=i)
            c_sheet = ws.cell(row=row_idx, column=3, value=item["sheet_name"])
            ws.cell(row=row_idx, column=4, value=item["filename"])
            c_rows = ws.cell(row=row_idx, column=5, value=item["rows"])
            c_cols = ws.cell(row=row_idx, column=6, value=item["cols"])
            c_link = ws.cell(
                row=row_idx, column=7, value=f'=HYPERLINK("#\'{item["sheet_name"]}\'!A1", "👉 Ir a Pestaña")'
            )

            c_num.alignment = Alignment(horizontal="center")
            c_sheet.font = Font(name="Segoe UI", bold=True, color="1E293B")
            c_rows.alignment = Alignment(horizontal="right")
            c_rows.number_format = "#,##0"
            c_cols.alignment = Alignment(horizontal="right")
            c_cols.number_format = "#,##0"
            c_link.font = Font(name="Segoe UI", bold=True, color=theme["header_fill"], underline="single")
            c_link.alignment = Alignment(horizontal="center")

            for col_idx in range(2, 8):
                cell = ws.cell(row=row_idx, column=col_idx)
                cell.border = border_all
                if row_fill and col_idx != 7:
                    cell.fill = row_fill

        ws.column_dimensions["A"].width = 3
        ws.column_dimensions["B"].width = 6
        ws.column_dimensions["C"].width = 25
        ws.column_dimensions["D"].width = 32
        ws.column_dimensions["E"].width = 14
        ws.column_dimensions["F"].width = 16
        ws.column_dimensions["G"].width = 18

    def _format_worksheet(
        self,
        ws: openpyxl.worksheet.worksheet.Worksheet,
        df: pd.DataFrame,
        theme: Dict[str, str],
        preserve_leading_zeros: bool = True,
        add_totals_row: bool = True,
        auto_fit_columns: bool = True,
        freeze_header: bool = True,
        sheet_index: int = 1,
    ) -> bool:
        """
        Aplica los estilos ejecutivos nativos de Excel (Ctrl + T), auto-filtros, formatos de número/texto y totales.
        """
        ws.views.sheetView[0].showGridLines = True

        header_fill = PatternFill(fgColor=theme["header_fill"], fill_type="solid")
        header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
        zebra_fill = PatternFill(fgColor=theme["zebra_fill"], fill_type="solid")

        border_side = Side(style="thin", color=theme["border_color"])
        cell_border = Border(left=border_side, right=border_side, top=border_side, bottom=border_side)

        total_top_border = Side(style="thin", color="1E293B")
        total_bottom_border = Side(style="double", color="1E293B")
        total_border = Border(left=border_side, right=border_side, top=total_top_border, bottom=total_bottom_border)

        # 1. Escribir Encabezados
        for col_idx, col_name in enumerate(df.columns, start=1):
            cell = ws.cell(row=1, column=col_idx, value=col_name)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = cell_border

        ws.row_dimensions[1].height = 28
        numeric_cols_idx = []

        # 2. Escribir Filas de Datos
        for row_idx, row_data in enumerate(df.itertuples(index=False), start=2):
            is_even = row_idx % 2 == 0
            ws.row_dimensions[row_idx].height = 20

            for col_idx, raw_val in enumerate(row_data, start=1):
                cell = ws.cell(row=row_idx, column=col_idx)
                val_str = str(raw_val).strip() if pd.notna(raw_val) else ""

                col_name = df.columns[col_idx - 1]
                col_type = self._infer_column_type(col_name, [val_str])

                if preserve_leading_zeros and (
                    col_type == "text_id"
                    or (val_str.startswith("0") and len(val_str) > 1 and val_str.isdigit())
                    or ("-" in val_str and val_str.replace("-", "").isdigit())
                ):
                    cell.value = val_str
                    cell.number_format = "@"
                    cell.alignment = Alignment(horizontal="center", vertical="center")

                elif col_type in ["currency", "number"] and val_str != "":
                    clean_num_str = re.sub(r"[^\d.-]", "", val_str.replace(",", "."))
                    try:
                        num_val = float(clean_num_str) if "." in clean_num_str else int(clean_num_str)
                        cell.value = num_val
                        if col_type == "currency":
                            cell.number_format = '"S/" #,##0.00'
                        else:
                            cell.number_format = "#,##0.00" if isinstance(num_val, float) else "#,##0"
                        cell.alignment = Alignment(horizontal="right", vertical="center")
                        if col_idx not in numeric_cols_idx:
                            numeric_cols_idx.append(col_idx)
                    except ValueError:
                        cell.value = val_str
                        cell.alignment = Alignment(horizontal="left", vertical="center")
                else:
                    cell.value = val_str
                    align_h = "center" if col_type == "date" else "left"
                    cell.alignment = Alignment(horizontal=align_h, vertical="center")

                cell.border = cell_border
                if is_even:
                    cell.fill = zebra_fill

        last_data_row = max(len(df) + 1, 2)

        # 3. Crear TABLA NATIVA DE EXCEL (Ctrl + T)
        if len(df) > 0:
            last_col_letter = get_column_letter(len(df.columns))
            table_ref = f"A1:{last_col_letter}{last_data_row}"

            clean_sheet_name = re.sub(r"\W+", "_", ws.title).strip("_")
            if not clean_sheet_name:
                clean_sheet_name = "Tabla"
            table_displayName = f"Tabla_{clean_sheet_name}_{sheet_index}"[:240]

            tab = Table(displayName=table_displayName, ref=table_ref)
            tab.tableStyleInfo = TableStyleInfo(
                name=theme.get("table_style", "TableStyleMedium9"),
                showFirstColumn=False,
                showLastColumn=False,
                showRowStripes=True,
                showColumnStripes=False,
            )
            ws.add_table(tab)

        # 4. Fila de Totales si corresponde
        if add_totals_row and len(df) > 0 and len(numeric_cols_idx) > 0:
            totals_row_idx = last_data_row + 1
            ws.row_dimensions[totals_row_idx].height = 24

            first_cell = ws.cell(row=totals_row_idx, column=1, value="TOTALES")
            first_cell.font = Font(name="Segoe UI", size=11, bold=True, color="1E293B")
            first_cell.alignment = Alignment(horizontal="left", vertical="center")
            first_cell.border = total_border

            for col_idx in range(1, len(df.columns) + 1):
                cell = ws.cell(row=totals_row_idx, column=col_idx)
                cell.border = total_border
                cell.font = Font(name="Segoe UI", size=11, bold=True, color="1E293B")

                if col_idx in numeric_cols_idx:
                    col_letter = get_column_letter(col_idx)
                    cell.value = f"=SUM({col_letter}2:{col_letter}{last_data_row})"
                    cell.number_format = '"S/" #,##0.00'
                    cell.alignment = Alignment(horizontal="right", vertical="center")

        # 5. Inmovilizar Fila 1
        if freeze_header:
            ws.freeze_panes = "A2"

        # 6. Auto-ajustar Ancho de Columnas
        if auto_fit_columns:
            for col in ws.columns:
                max_len = 0
                col_letter = get_column_letter(col[0].column)
                for cell in col:
                    val = str(cell.value or "")
                    if val.startswith("="):
                        val = "S/ 999,999.00"
                    max_len = max(max_len, len(val))
                ws.column_dimensions[col_letter].width = min(max(max_len + 4, 12), 48)

        return True
